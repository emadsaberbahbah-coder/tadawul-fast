#!/usr/bin/env python3
"""
tfb_export_audit.py - offline six-gate audit of a TFB workbook export.

WHY (v1.0.0, 2026-10-03)
------------------------
Every TFB export is supposed to pass the six-gate audit: (a) data integrity,
(b) recommendation coherence, (c) cross-version consistency, (d) forecast vs
reality, (e) guard health from _Run_Log, (f) verdict + fix list. Until now the
gates were run as ad-hoc scripts, so two sessions on the same export could
disagree, and the in-workbook validator (Dashboard_Audit) samples only 1,500
rows per page (P-190). This script runs the full population of every tab in
one pass and prints the same numbers every time.

It is READ-ONLY and STANDALONE: it needs openpyxl and nothing from the repo,
so it runs on any machine against any .xlsx export of "Market Share
Deepseek-V3". It never opens the live Sheet, never writes to it, and never
calls a provider. Sharing / ACL state cannot be checked offline and is
reported as NOT CHECKED, never as PASS.

Checks implemented (first delivery, 2026-10-03 reconciliation):
  (a) page row counts vs expected, in-page and cross-page duplicate symbols,
      blank symbols, missing names, invalid prices, |day move| > 25 %, price
      outside the 52-week band, freshness by Last Updated (UTC), stuck rows,
      schema width 115 / 122, profit-margin unit errors (> 1.0 under a
      percent format), canned forecast-confidence concentration, crypto
      wrong-instrument names (curated map), zero forecast prices at positive
      current prices, forecast-price vs Expected-ROI pair disagreement
      (rounding-aware tolerance), Horizon Days vs Invest Period Label,
      _PIT_Fundamentals timestamps stored as reliability.
  (b) INVEST rows that fail the cockpit's DQ / reliability screen, sector-cap
      panel values vs policy, SELL-class rows carrying INVEST.
  (c) version stamps gathered from every surface, optionally compared with a
      manifest JSON.
  (d) Performance_Log matured win rate, duplicate matured keys, all-zero
      risk columns; Signal_History multi-version symbol-days; calibration
      state; Brier score from _Status; S-1 gate verdict and a criterion
      that passes on an empty _Corporate_Actions table; Hypothesis_Registry
      verdicts.
  (e) _Run_Log ERROR / WARNING counts, EODHD daily usage vs target, HTTP 402
      rows, identity-guard refusal rows, run start lateness vs the cron
      slots, the cockpit-vs-sync race (Top 10 run time inside the sync run).
  Book: ledger active lots x My_Portfolio prices x FX vs Portfolio_Decision
      KPI (sold holdings still displayed, stale KPI cash), _Cash_Snapshot
      duplicate dates and stale notes, fills vs _Trade_Notes.
  (f) RAG per lane (DATA, BOOK, DECISION, PIPELINE, MODEL, GOVERNANCE) and
      an ordered fix list.

Usage
-----
  python scripts/tfb_export_audit.py export.xlsx [--json out.json] [--md out.md]
        [--expect Market_Leaders=255,Global_Markets=6609,...]
        [--dq 80] [--rel 70] [--sector-cap 3/40] [--eodhd-target 90000]
        [--cron "17 4,12,20"] [--tz-offset 3] [--since-days 7]
        [--manifest versions.json] [--asof 2026-10-03]
  python scripts/tfb_export_audit.py --selftest

Exit code: 0 = no FAIL, 1 = at least one FAIL, 2 = usage / load error.
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import hashlib
import json
import os
import re
import sys
import tempfile

SCRIPT_VERSION = "1.0.0"

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.stderr.write("openpyxl is required: pip install openpyxl\n")
    sys.exit(2)

MARKET_PAGES = ["Market_Leaders", "Global_Markets", "Commodities_FX", "Mutual_Funds"]
DEFAULT_EXPECT = {"Market_Leaders": 255, "Global_Markets": 6609,
                  "Commodities_FX": 453, "Mutual_Funds": 2474}
MAIN_WIDTH = 115
PORTFOLIO_WIDTH = 122
FX_TO_SAR = {"SAR": 1.0, "USD": 3.75}
DEAD_TABS = ["_Diag_Response_Shape", "_Data_Audit", "Diagnostic_Report", "Sheet97",
             "Sheet88", "_SM_TEST", "_Shariah_Upload", "_SM_TEST_LIVE", "System_Logs"]
LABEL_DAYS = {"1W": 7, "2W": 14, "1M": 30, "3M": 90, "6M": 180, "12M": 365, "1Y": 365}
SELL_CLASS = {"SELL", "STRONG_SELL", "REDUCE", "AVOID"}

# Curated crypto identity map: base symbol -> name keywords (any must appear).
CRYPTO_NAMES = {
    "BTC": ["bitcoin"], "ETH": ["ethereum"], "SOL": ["solana"], "XRP": ["xrp", "ripple"],
    "BNB": ["bnb", "binance"], "ADA": ["cardano"], "DOGE": ["dogecoin"], "TRX": ["tron"],
    "AVAX": ["avalanche"], "DOT": ["polkadot"], "LINK": ["chainlink"], "MATIC": ["polygon"],
    "POL": ["polygon"], "LTC": ["litecoin"], "SHIB": ["shiba"], "UNI": ["uniswap"],
    "APT": ["aptos"], "SUI": ["sui"], "COMP": ["compound"], "ATOM": ["cosmos"],
    "NEAR": ["near"], "ARB": ["arbitrum"], "OP": ["optimism"], "AAVE": ["aave"],
    "MKR": ["maker"], "LDO": ["lido"], "TIA": ["celestia"], "EGLD": ["multiversx", "elrond"],
    "FIL": ["filecoin"], "ICP": ["internet computer"], "XLM": ["stellar"],
    "BCH": ["bitcoin cash"], "ETC": ["ethereum classic"], "HBAR": ["hedera"],
    "VET": ["vechain"], "ALGO": ["algorand"], "IMX": ["immutable"], "INJ": ["injective"],
    "RNDR": ["render"], "RENDER": ["render"], "GRT": ["graph"], "SAND": ["sandbox"],
    "MANA": ["decentraland"], "AXS": ["axie"], "FTM": ["fantom"], "XMR": ["monero"],
    "TON": ["toncoin", "ton"], "PEPE": ["pepe"], "SEI": ["sei"], "STX": ["stacks"],
    "KAS": ["kaspa"], "TAO": ["bittensor"], "USDT": ["tether"], "USDC": ["usd coin", "usdc"],
    "DAI": ["dai"], "CRO": ["cronos", "crypto.com"], "QNT": ["quant"], "THETA": ["theta"],
    "FLOW": ["flow"], "CHZ": ["chiliz"], "ENS": ["ethereum name"], "GALA": ["gala"],
    "APE": ["apecoin"], "CRV": ["curve"], "SNX": ["synthetix"], "RUNE": ["thorchain"],
    "KAVA": ["kava"], "ZEC": ["zcash"], "DASH": ["dash"], "NEO": ["neo"], "EOS": ["eos"],
    "XTZ": ["tezos"], "MINA": ["mina"], "ROSE": ["oasis"], "ONE": ["harmony"],
    "ZIL": ["zilliqa"], "ENJ": ["enjin"], "BAT": ["basic attention"], "1INCH": ["1inch"],
    "SUSHI": ["sushi"], "YFI": ["yearn"], "DYDX": ["dydx"], "GMX": ["gmx"],
    "PENDLE": ["pendle"], "JUP": ["jupiter"], "PYTH": ["pyth"], "ONDO": ["ondo"],
    "ENA": ["ethena"], "BONK": ["bonk"], "FLOKI": ["floki"], "AR": ["arweave"],
    "FET": ["fetch", "artificial superintelligence"], "WLD": ["worldcoin"],
    "BLUR": ["blur"], "CFX": ["conflux"], "MNT": ["mantle"], "STRK": ["starknet"],
    "AKT": ["akash"], "JASMY": ["jasmy"], "IOTA": ["iota"], "HNT": ["helium"],
    "RAY": ["raydium"], "CAKE": ["pancake"], "LUNC": ["terra"], "LUNA": ["terra"],
    "FXS": ["frax"], "BTT": ["bittorrent"], "NEXO": ["nexo"], "GNO": ["gnosis"],
    "RPL": ["rocket pool"], "CVX": ["convex"], "BAL": ["balancer"], "LRC": ["loopring"],
    "ANKR": ["ankr"], "CELO": ["celo"], "KSM": ["kusama"], "GLMR": ["moonbeam"],
    "ASTR": ["astar"], "WAVES": ["waves"], "ZRX": ["0x"], "QTUM": ["qtum"], "ICX": ["icon"],
    "ONT": ["ontology"], "STORJ": ["storj"], "SC": ["siacoin"], "DCR": ["decred"],
    "RVN": ["ravencoin"], "XEC": ["ecash"], "BSV": ["bitcoin sv"], "XEM": ["nem"],
    "LSK": ["lisk"], "DGB": ["digibyte"], "XNO": ["nano"], "HOT": ["holo"],
    "CKB": ["nervos"], "MASK": ["mask"], "LPT": ["livepeer"], "UMA": ["uma"],
    "KNC": ["kyber"], "PAXG": ["pax gold"], "XAUT": ["tether gold"], "TUSD": ["trueusd"],
    "FDUSD": ["first digital"], "PYUSD": ["paypal"], "WIF": ["dogwifhat", "wif"],
    "NOT": ["notcoin"], "ORDI": ["ordi"], "BEAM": ["beam"], "W": ["wormhole"],
    "ETHFI": ["ether.fi"], "AERO": ["aerodrome"], "ZK": ["zksync"], "TWT": ["trust wallet"],
}


# --------------------------------------------------------------------------- #
# Small helpers
# --------------------------------------------------------------------------- #
def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def _s(v):
    return "" if v is None else str(v).strip()


def _parse_dt(v, tz_offset_hours=3):
    """Datetime / ISO string -> aware UTC datetime, or None.
    Naive values are taken as workbook-local (Riyadh by default)."""
    local = dt.timezone(dt.timedelta(hours=tz_offset_hours))
    if isinstance(v, dt.datetime):
        d = v
    elif isinstance(v, dt.date):
        d = dt.datetime(v.year, v.month, v.day)
    else:
        s = _s(v)
        if not s:
            return None
        s = s.replace("Z", "+00:00")
        try:
            d = dt.datetime.fromisoformat(s)
        except ValueError:
            m = re.search(r"(\d{4}-\d{2}-\d{2})[ T](\d{2}:\d{2}(?::\d{2})?)", s)
            if not m:
                m2 = re.search(r"(\d{4}-\d{2}-\d{2})", s)
                if not m2:
                    return None
                d = dt.datetime.fromisoformat(m2.group(1))
            else:
                d = dt.datetime.fromisoformat(m.group(1) + "T" + m.group(2))
    if d.tzinfo is None:
        d = d.replace(tzinfo=local)
    return d.astimezone(dt.timezone.utc)


def _date_of(v):
    """Date part (YYYY-MM-DD) of a cell that holds a date/datetime/string."""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    m = re.search(r"\d{4}-\d{2}-\d{2}", _s(v))
    return m.group(0) if m else None


def _looks_like_timestamp(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return True
    return bool(re.match(r"^\d{4}-\d{2}-\d{2}([ T]\d{2}:\d{2})?", _s(v)))


def _is_crypto(sym):
    return bool(re.match(r"^[A-Z0-9]+-USD$", sym or ""))


def _panel_value(rows, label):
    """Value to the right of a panel label such as 'T10: Max Per Sector'."""
    for r in rows:
        for i, c in enumerate(r):
            if _s(c) == label and i + 1 < len(r):
                return r[i + 1]
    return None


def _sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


class Finding(dict):
    """gate, lane, check, status (PASS|WARN|FAIL|INFO|NOT_CHECKED), count, detail, examples."""


def F(gate, lane, check, status, count=None, detail="", examples=None):
    return Finding(gate=gate, lane=lane, check=check, status=status, count=count,
                   detail=detail, examples=list(examples or [])[:8])


# --------------------------------------------------------------------------- #
# Workbook access (one read, trailing-empty rows trimmed)
# --------------------------------------------------------------------------- #
class Book:
    def __init__(self, path):
        self.path = path
        self.sheets = {}      # title -> list of row tuples
        self.state = {}       # title -> 'visible' | 'hidden'
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
            while rows and all(v is None or v == "" for v in rows[-1]):
                rows.pop()
            # read_only mode returns ragged tuples; pad every row to the sheet width so
            # column indexes taken from the header are always in range
            width = max((len(r) for r in rows), default=0)
            rows = [r + (None,) * (width - len(r)) if len(r) < width else r for r in rows]
            self.sheets[ws.title] = rows
            self.state[ws.title] = ws.sheet_state
        wb.close()

    def rows(self, name):
        return self.sheets.get(name)

    def table(self, name, header_pred=None, header_row=0):
        """(index-by-header, body rows). header_pred(row) picks the header row."""
        rows = self.rows(name)
        if not rows:
            return {}, []
        hi = header_row
        if header_pred is not None:
            hi = next((i for i, r in enumerate(rows) if header_pred(r)), None)
            if hi is None:
                return {}, []
        header = [_s(h) for h in rows[hi]]
        ix = {}
        for i, h in enumerate(header):
            if h and h not in ix:
                ix[h] = i
        body = [r for r in rows[hi + 1:] if any(v not in (None, "") for v in r)]
        return ix, body


# --------------------------------------------------------------------------- #
# Gate (a): structure and market pages
# --------------------------------------------------------------------------- #
def check_structure(book, now_utc):
    out = []
    names = list(book.sheets)
    vis = [n for n in names if book.state.get(n) == "visible"]
    out.append(F("a", "DATA", "tabs", "INFO", len(names),
                 f"{len(vis)} visible, {len(names) - len(vis)} hidden"))
    dead = [n for n in DEAD_TABS if n in book.sheets]
    out.append(F("a", "GOVERNANCE", "dead_tabs_present", "WARN" if dead else "PASS",
                 len(dead), "retired or test tabs still in the workbook", dead))
    rows = book.rows("_Decision_Diagnostics") or []
    stamp = next((_parse_dt(m.group(0)) for r in rows[:3] for c in r
                  if (m := re.search(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}", _s(c)))), None)
    if stamp:
        age = (now_utc - stamp).days
        out.append(F("a", "GOVERNANCE", "decision_diagnostics_age_days",
                     "WARN" if age > 7 else "PASS", age, f"last ran {stamp:%Y-%m-%d}"))
    da = book.rows("_Data_Audit") or []
    errs = sum(1 for r in da for c in r if _s(c).startswith("#") and _s(c).endswith("!"))
    if da:
        out.append(F("a", "GOVERNANCE", "data_audit_error_cells", "WARN" if errs else "PASS",
                     errs, "formula error literals in _Data_Audit"))
    return out


def check_market_pages(book, expect, asof_date, dq_floor, rel_floor):
    out, metrics, allsyms = [], {}, collections.Counter()
    invest_total = invest_pass = 0
    for page in MARKET_PAGES:
        ix, body = book.table(page)
        if not ix:
            out.append(F("a", "DATA", f"{page}.present", "FAIL", 0, "page missing"))
            continue
        width = max(len(r) for r in (book.rows(page) or [()]))
        hdr_width = len([h for h in book.rows(page)[0] if h not in (None, "")])
        m = {"rows": len(body), "header_width": hdr_width, "width": width}
        exp = expect.get(page)
        out.append(F("a", "DATA", f"{page}.rows",
                     "PASS" if exp is None or len(body) == exp else "FAIL", len(body),
                     f"expected {exp}" if exp is not None else "no expectation set"))
        out.append(F("a", "DATA", f"{page}.schema_width",
                     "PASS" if hdr_width == MAIN_WIDTH else "FAIL", hdr_width,
                     f"contract {MAIN_WIDTH} columns"))
        sym_i, name_i = ix.get("Symbol"), ix.get("Name")
        price_i, prev_i = ix.get("Current Price"), ix.get("Previous Close")
        hi_i, lo_i = ix.get("52W High"), ix.get("52W Low")
        upd_i = ix.get("Last Updated (UTC)")
        pm_i, gm_i, om_i = ix.get("Profit Margin"), ix.get("Gross Margin"), ix.get("Operating Margin")
        conf_i = ix.get("Forecast Confidence")
        fa_i, dq_i, rel_i = ix.get("Final Action"), ix.get("Data Quality Score"), ix.get("Forecast Reliability Score")
        rec_i = ix.get("Recommendation")

        syms = [_s(r[sym_i]) if sym_i is not None else "" for r in body]
        c = collections.Counter(s for s in syms if s)
        dup = sum(v - 1 for v in c.values() if v > 1)
        allsyms.update(set(s for s in syms if s))
        blank_sym = sum(1 for s in syms if not s)
        miss_name = [syms[i] for i, r in enumerate(body) if name_i is not None and not _s(r[name_i])]
        no_price = [syms[i] for i, r in enumerate(body)
                    if price_i is not None and (_num(r[price_i]) is None or _num(r[price_i]) <= 0)]
        big = []
        out52 = []
        for i, r in enumerate(body):
            p = _num(r[price_i]) if price_i is not None else None
            pc = _num(r[prev_i]) if prev_i is not None else None
            if p and pc and pc > 0 and abs(p / pc - 1) > 0.25:
                big.append(f"{syms[i]} {p:g}/{pc:g}")
            h = _num(r[hi_i]) if hi_i is not None else None
            lo = _num(r[lo_i]) if lo_i is not None else None
            if p and h and lo and (p > h * 1.01 or p < lo * 0.99):
                out52.append(f"{syms[i]} {p:g} vs {lo:g}-{h:g}")
        dates = collections.Counter(_date_of(r[upd_i]) for r in body if upd_i is not None)
        no_date = dates.pop(None, 0)
        fresh = dates.get(asof_date, 0) if asof_date else None
        stale = []
        if asof_date and upd_i is not None:
            cutoff = dt.date.fromisoformat(asof_date) - dt.timedelta(days=5)
            for i, r in enumerate(body):
                d = _date_of(r[upd_i])
                if d and dt.date.fromisoformat(d) < cutoff:
                    stale.append(f"{syms[i]} {d}")
        pm_vals = [_num(r[pm_i]) for r in body if pm_i is not None]
        pm_n = sum(1 for v in pm_vals if v is not None)
        pm_bad = sum(1 for v in pm_vals if v is not None and abs(v) > 1.0)
        gm_bad = sum(1 for r in body if gm_i is not None and (_num(r[gm_i]) or 0) > 1.0)
        om_bad = sum(1 for r in body if om_i is not None and (_num(r[om_i]) or 0) > 1.0)
        confs = collections.Counter(_num(r[conf_i]) for r in body if conf_i is not None and _num(r[conf_i]) is not None)
        top_conf, top_n = (confs.most_common(1)[0] if confs else (None, 0))
        conf_share = (top_n / len(body)) if body else 0
        inv = inv_ok = sell_inv = 0
        for r in body:
            if fa_i is not None and _s(r[fa_i]).upper() == "INVEST":
                inv += 1
                d_, rl = (_num(r[dq_i]) if dq_i is not None else None), (_num(r[rel_i]) if rel_i is not None else None)
                if d_ is not None and rl is not None and d_ >= dq_floor and rl >= rel_floor:
                    inv_ok += 1
                if rec_i is not None and _s(r[rec_i]).upper() in SELL_CLASS:
                    sell_inv += 1
        invest_total += inv
        invest_pass += inv_ok
        m.update(dup_symbols=dup, blank_symbols=blank_sym, missing_names=len(miss_name),
                 invalid_price=len(no_price), big_moves=len(big), outside_52w=len(out52),
                 fresh_rows=fresh, stale_rows=len(stale), pm_over_1=pm_bad, pm_n=pm_n,
                 gm_over_1=gm_bad, om_over_1=om_bad, top_conf=top_conf, top_conf_rows=top_n,
                 invest=inv, invest_eligible=inv_ok, sell_class_invest=sell_inv,
                 dates=dict(dates.most_common(5)), no_update_date=no_date)
        metrics[page] = m
        out.append(F("a", "DATA", f"{page}.duplicate_symbols", "PASS" if dup == 0 else "FAIL", dup))
        out.append(F("a", "DATA", f"{page}.blank_symbols", "PASS" if blank_sym == 0 else "FAIL", blank_sym))
        out.append(F("a", "DATA", f"{page}.missing_names", "PASS" if not miss_name else "WARN",
                     len(miss_name), "rows with no Name", miss_name))
        out.append(F("a", "DATA", f"{page}.invalid_price", "PASS" if not no_price else "WARN",
                     len(no_price), "Current Price blank or <= 0", no_price))
        out.append(F("a", "DATA", f"{page}.day_move_over_25pct", "PASS" if not big else "WARN",
                     len(big), "|price/prev close - 1| > 25 %; review", big))
        out.append(F("a", "DATA", f"{page}.price_outside_52w", "PASS" if not out52 else "WARN",
                     len(out52), "price beyond the 52-week band by > 1 %", out52))
        if asof_date and upd_i is not None:
            cov = fresh / len(body) if body else 0
            out.append(F("a", "DATA", f"{page}.fresh_on_asof", "PASS" if cov >= 0.95 else "WARN",
                         fresh, f"{cov:.1%} of rows carry Last Updated (UTC) = {asof_date}"))
            out.append(F("a", "DATA", f"{page}.stale_over_5_days", "PASS" if not stale else "WARN",
                         len(stale), "Last Updated older than 5 days", stale))
        if pm_n:
            share = pm_bad / pm_n
            out.append(F("a", "DATA", f"{page}.profit_margin_unit_errors",
                         "PASS" if share < 0.05 else "FAIL", pm_bad,
                         f"{share:.1%} of {pm_n} populated margins are > 1.0 (P-152 unit error)"))
        if gm_bad or om_bad:
            out.append(F("a", "DATA", f"{page}.gross_operating_margin_over_1", "WARN",
                         gm_bad + om_bad, f"gross {gm_bad}, operating {om_bad}"))
        if confs:
            out.append(F("a", "DATA", f"{page}.canned_confidence", "PASS" if conf_share < 0.25 else "WARN",
                         top_n, f"value {top_conf} on {conf_share:.1%} of rows"))
        if sell_inv:
            out.append(F("b", "DECISION", f"{page}.sell_class_rows_marked_invest", "FAIL", sell_inv,
                         "Recommendation in SELL/REDUCE/AVOID with Final Action INVEST"))
    cross = sorted(s for s, v in allsyms.items() if v > 1)
    out.append(F("a", "DATA", "cross_page_duplicate_symbols", "PASS" if not cross else "WARN",
                 len(cross), "symbol present on more than one page (cockpit de-duplicates)", cross))
    out.append(F("b", "DECISION", "invest_rows_vs_eligibility",
                 "PASS" if invest_total == invest_pass else "FAIL", invest_total - invest_pass,
                 f"{invest_total} rows say INVEST; {invest_pass} meet DQ >= {dq_floor} and "
                 f"reliability >= {rel_floor} (label vs executable eligibility)"))
    metrics["invest_total"], metrics["invest_eligible"] = invest_total, invest_pass
    return out, metrics


def check_crypto_identity(book):
    out = []
    ix, body = book.table("Commodities_FX")
    if not ix:
        return out
    sym_i, name_i, price_i = ix.get("Symbol"), ix.get("Name"), ix.get("Current Price")
    f_cols = [ix[k] for k in ("Forecast Price 1M", "Forecast Price 3M", "Forecast Price 12M") if k in ix]
    mism, unmapped, zero_fc = [], 0, []
    for r in body:
        sym = _s(r[sym_i]) if sym_i is not None else ""
        if not _is_crypto(sym):
            continue
        base = sym[:-4]
        name = _s(r[name_i]).lower() if name_i is not None else ""
        keys = CRYPTO_NAMES.get(base)
        if keys is None:
            unmapped += 1
        elif not any(k in name for k in keys):
            mism.append(f"{sym} = '{_s(r[name_i])}'")
        p = _num(r[price_i]) if price_i is not None else None
        fvals = [_num(r[i]) for i in f_cols]
        # a blank forecast is "not computed"; only a numeric zero at a positive price is precision loss
        if p and p > 0 and any(v is not None and v <= 0 for v in fvals):
            zero_fc.append(f"{sym} price {p:g}")
    out.append(F("a", "DATA", "crypto_wrong_instrument_name", "PASS" if not mism else "FAIL",
                 len(mism), f"name does not match the curated map ({unmapped} symbols unmapped, not judged)", mism))
    out.append(F("a", "DATA", "crypto_zero_forecast_price", "PASS" if not zero_fc else "FAIL",
                 len(zero_fc), "forecast price 0 with a positive current price (precision loss)", zero_fc))
    return out


def check_forecast_pairs(book):
    """Forecast Price 12M vs Expected ROI 12M, rounding-aware."""
    out, total, bad_all = [], 0, []
    for page in MARKET_PAGES:
        ix, body = book.table(page)
        if not ix or "Forecast Price 12M" not in ix or "Expected ROI 12M" not in ix:
            continue
        sym_i, p_i, f_i, roi_i = ix.get("Symbol"), ix.get("Current Price"), ix["Forecast Price 12M"], ix["Expected ROI 12M"]
        bad = []
        for r in body:
            p, f, roi = _num(r[p_i]), _num(r[f_i]), _num(r[roi_i])
            if not p or p <= 0 or f is None or f <= 0 or roi is None:
                continue
            total += 1
            implied = f / p - 1
            tol = max(0.001, 0.00005 / p * 1.02)   # 0.1 pp, or 4-decimal rounding of the forecast price
            if abs(implied - roi) > tol:
                bad.append(f"{_s(r[sym_i])} f12 {f:g} @ {p:g} -> {implied:+.2%} vs ROI {roi:+.2%}")
        bad_all += bad
        if bad:
            out.append(F("a", "DATA", f"{page}.forecast_roi_pair_mismatch", "FAIL", len(bad),
                         "Forecast Price 12M and Expected ROI 12M disagree beyond rounding (P-158)", bad))
    out.append(F("a", "DATA", "forecast_roi_pairs_total", "PASS" if not bad_all else "FAIL",
                 len(bad_all), f"{len(bad_all)} of {total} pairs disagree"))
    return out


def check_horizon(book):
    out = []
    for page in ["My_Portfolio"] + MARKET_PAGES:
        ix, body = book.table(page)
        if not ix or "Horizon Days" not in ix or "Invest Period Label" not in ix:
            continue
        bad = []
        for r in body:
            days, label = _num(r[ix["Horizon Days"]]), _s(r[ix["Invest Period Label"]]).upper()
            want = LABEL_DAYS.get(label)
            if days is not None and want is not None and abs(days - want) > want * 0.2:
                bad.append(f"{_s(r[ix.get('Symbol', 0)])} {days:g}d vs {label}")
        if page == "My_Portfolio" or bad:
            systematic = body and len(bad) / len(body) > 0.5
            out.append(F("a", "DATA", f"{page}.horizon_vs_label",
                         "PASS" if not bad else ("WARN" if systematic else "FAIL"), len(bad),
                         ("systematic: Horizon Days and Invest Period Label follow different conventions on "
                          f"{len(bad)} of {len(body)} rows; define which one drives exits (P-194)") if systematic
                         else "Horizon Days disagrees with Invest Period Label by > 20 %", bad))
    return out


def check_pit(book):
    ix, body = book.table("_PIT_Fundamentals")
    if not ix or "Forecast Reliability" not in ix:
        return []
    bad = [r for r in body if _looks_like_timestamp(r[ix["Forecast Reliability"]])]
    return [F("a", "MODEL", "pit_timestamp_in_reliability", "PASS" if not bad else "FAIL", len(bad),
              f"of {len(body)} rows hold a date/time where a reliability score belongs",
              [_s(r[ix.get('Symbol', 0)]) + " " + _s(r[ix["Forecast Reliability"]])[:19] for r in bad])]


# --------------------------------------------------------------------------- #
# Book: ledger, cash, Portfolio_Decision
# --------------------------------------------------------------------------- #
def _status_line(rows, row=1, col=1):
    try:
        return _s(rows[row][col])
    except (IndexError, TypeError):
        return ""


def check_portfolio(book, tz_offset, since_days, now_utc):
    out, metrics = [], {}
    # Ledger
    lrows = book.rows("_Portfolio_CostBasis") or []
    lix, lbody = book.table("_Portfolio_CostBasis",
                            header_pred=lambda r: _s(r[0]) == "Symbol" and "Status" in [_s(c) for c in r])
    active, fills = [], 0
    since = now_utc - dt.timedelta(days=since_days)
    for r in lbody:
        st = _s(r[lix.get("Status", 3)])
        sym, ccy = _s(r[lix.get("Symbol", 0)]), _s(r[lix.get("Ccy", 2)])
        shares, buy = _num(r[lix.get("Shares", 6)]), _num(r[lix.get("Buy Price", 5)])
        bd, sd = _parse_dt(r[lix.get("Buy Date", 4)], tz_offset), _parse_dt(r[lix.get("Sell Date", 8)], tz_offset)
        if bd and bd >= since:
            fills += 1
        if sd and sd >= since:
            fills += 1
        if st.lower() == "active":
            active.append((sym, ccy, shares or 0, buy or 0))
    status = _status_line(lrows)
    m = re.search(r"closed (-?[\d,]+) · active (-?[\d,]+) · lifetime (-?[\d,]+)", status)
    if m:
        metrics["ledger_totals"] = {k: float(v.replace(",", "")) for k, v in
                                    zip(("closed", "active", "lifetime"), m.groups())}
    # Prices from My_Portfolio
    pix, pbody = book.table("My_Portfolio")
    prices = {_s(r[pix["Symbol"]]): _num(r[pix["Current Price"]]) for r in pbody} if pix else {}
    if pix:
        w = len([h for h in book.rows("My_Portfolio")[0] if h not in (None, "")])
        out.append(F("a", "DATA", "My_Portfolio.schema_width", "PASS" if w == PORTFOLIO_WIDTH else "FAIL",
                     w, f"contract {PORTFOLIO_WIDTH} columns"))
    holdings_sar, unpriced = 0.0, []
    for sym, ccy, sh, _ in active:
        fx, p = FX_TO_SAR.get(ccy), prices.get(sym)
        if fx is None or p is None:
            unpriced.append(sym)
            continue
        holdings_sar += sh * p * fx
    # Cash
    cix, cbody = book.table("_Cash_Snapshot")
    cash = None
    dup_dates = 0
    stale_notes = 0
    if cix and "Balance SAR" in cix:
        bals = [(r[cix.get("Date", 0)], _num(r[cix["Balance SAR"]]), _s(r[cix.get("Note", 6)]))
                for r in cbody if _num(r[cix["Balance SAR"]]) is not None]
        if bals:
            cash = bals[-1][1]
        dc = collections.Counter(_date_of(b[0]) for b in bals)
        dup_dates = sum(v - 1 for v in dc.values() if v > 1)
        for a, b in zip(bals, bals[1:]):
            if a[1] != b[1] and a[2] and a[2] == b[2]:
                stale_notes += 1
        metrics["cash_sar"] = cash
    # Portfolio_Decision
    prow = book.rows("Portfolio_Decision") or []
    pd_status = _status_line(prow)
    kpi = {}
    for i, r in enumerate(prow):
        if _s(r[0]) == "Portfolio (SAR)" and i + 1 < len(prow):
            hdr = [_s(c) for c in r]
            kpi = {h: _num(v) for h, v in zip(hdr, prow[i + 1]) if h}
            break
    actions = []
    for i, r in enumerate(prow):
        if _s(r[0]) == "Action" and _s(r[1]) == "Symbol":
            hdr = [_s(c) for c in r]
            for rr in prow[i + 1:]:
                if not _s(rr[0]):
                    break
                actions.append({h: v for h, v in zip(hdr, rr) if h})
            break
    pd_cash_input = _num(_panel_value(prow, "PF: Cash Available SAR"))
    active_syms = {a[0] for a in active}
    pd_syms = [_s(a.get("Symbol")) for a in actions]
    sold_shown = [s for s in pd_syms if s and s not in active_syms]
    missing = sorted(active_syms - set(pd_syms))
    metrics.update(active_lots=len(active), holdings_sar=round(holdings_sar, 2),
                   nav_sar=(round(holdings_sar + cash, 2) if cash is not None else None),
                   pd_kpi=kpi, pd_actions=len(actions), pd_cash_input=pd_cash_input,
                   fills_in_window=fills)
    out.append(F("book", "BOOK", "ledger_active_lots_priced", "PASS" if not unpriced else "WARN",
                 len(active) - len(unpriced), f"{len(active)} active lots; unpriced: {unpriced}"))
    out.append(F("book", "BOOK", "portfolio_decision_sold_holdings_shown", "PASS" if not sold_shown else "FAIL",
                 len(sold_shown), "symbols in Portfolio_Decision that are not Active in the ledger", sold_shown))
    if missing:
        out.append(F("book", "BOOK", "portfolio_decision_missing_active", "FAIL", len(missing),
                     "active ledger lots absent from Portfolio_Decision", missing))
    if cash is not None and kpi.get("Cash (SAR)") is not None:
        d = kpi["Cash (SAR)"] - cash
        out.append(F("book", "BOOK", "portfolio_decision_kpi_cash_vs_snapshot", "PASS" if abs(d) <= 1 else "FAIL",
                     round(d, 2), f"KPI cash {kpi['Cash (SAR)']:,.0f} vs latest _Cash_Snapshot {cash:,.2f} SAR"))
    if cash is not None and pd_cash_input is not None:
        d = pd_cash_input - cash
        out.append(F("book", "BOOK", "panel_cash_vs_snapshot", "PASS" if abs(d) <= 1 else "WARN",
                     round(d, 2), f"panel input {pd_cash_input:,.2f} vs snapshot {cash:,.2f}"))
    if cash is not None and kpi.get("Portfolio (SAR)") is not None and not unpriced:
        nav = holdings_sar + cash
        d = kpi["Portfolio (SAR)"] - nav
        out.append(F("book", "BOOK", "portfolio_decision_kpi_nav_vs_recomputed", "PASS" if abs(d) <= 100 else "WARN",
                     round(d, 2), f"KPI {kpi['Portfolio (SAR)']:,.0f} vs recomputed {nav:,.2f} "
                                  f"(holdings {holdings_sar:,.2f} + cash {cash:,.2f})"))
    out.append(F("book", "BOOK", "cash_snapshot_duplicate_dates", "PASS" if not dup_dates else "WARN", dup_dates))
    out.append(F("book", "BOOK", "cash_snapshot_stale_notes", "PASS" if not stale_notes else "WARN", stale_notes,
                 "balance changed while the note text stayed identical"))
    m = re.search(r"actions v([\d.]+)", pd_status)
    if m:
        metrics["portfolio_actions_version"] = m.group(1)
    m = re.search(r"Last run (\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})", pd_status)
    if m:
        metrics["pd_last_run_utc"] = _parse_dt(m.group(1), tz_offset).isoformat()
    # Trade notes
    tix, tbody = book.table("_Trade_Notes")
    notes = len([r for r in tbody if _s(r[0])]) if tix else 0
    metrics["trade_notes"] = notes
    out.append(F("gov", "GOVERNANCE", "trade_notes_vs_fills", "PASS" if notes >= fills else "FAIL",
                 fills - notes, f"{fills} fills in the last {since_days} days vs {notes} trade notes"))
    return out, metrics


# --------------------------------------------------------------------------- #
# Decision layer: Top 10 cockpit, timing race, caps
# --------------------------------------------------------------------------- #
def _status_pages(book, tz_offset):
    rows = book.rows("_Status") or []
    pages = {}
    globals_ = {}
    for r in rows:
        if _s(r[0]) in MARKET_PAGES:
            msg = _s(r[3]) if len(r) > 3 else ""
            pages[_s(r[0])] = {"updated_utc": _parse_dt(r[1], tz_offset), "status": _s(r[2]), "message": msg,
                               "run": (re.search(r"run=(\d+)", msg) or [None, None])[1],
                               "version": (re.search(r"\[STATUS-STAMP v([\d.]+)\]", msg) or [None, None])[1]}
        if len(r) > 12 and _s(r[11]):
            globals_[_s(r[11])] = _s(r[12])
    return pages, globals_


def check_cockpit(book, tz_offset, sector_cap):
    out, metrics = [], {}
    rows = book.rows("Top_10_Investments") or []
    st = _status_line(rows)
    m_run = re.search(r"Last run (\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})", st)
    t10 = _parse_dt(m_run.group(1), tz_offset) if m_run else None
    output = (re.search(r"output: (\w+)", st) or [None, None])[1]
    # the "FEED NOT ACTIONABLE - aged:GM" banner sits on its own row below the panel
    aged = sorted(set(re.findall(r"aged:(\w+)", " ".join(_s(c) for r in rows[:25] for c in r))))
    metrics.update(t10_last_run_utc=t10.isoformat() if t10 else None, t10_output=output,
                   t10_builder=(re.search(r"builder v([\d.]+)", st) or [None, None])[1],
                   t10_route=(re.search(r"route v([\d.]+)", st) or [None, None])[1])
    pages, _ = _status_pages(book, tz_offset)
    stamps = [p["updated_utc"] for p in pages.values() if p["updated_utc"]]
    if t10 and stamps:
        first, last = min(stamps), max(stamps)
        runs = {p["run"] for p in pages.values()}
        if len(runs) > 1:
            out.append(F("b", "DECISION", "status_pages_mixed_runs", "FAIL", len(runs),
                         "the four page stamps come from different sync runs", sorted(str(x) for x in runs)))
        if first <= t10 < last:
            out.append(F("b", "DECISION", "cockpit_ran_inside_sync_run", "FAIL",
                         int((last - t10).total_seconds()),
                         f"Top 10 ran {int((last - t10).total_seconds())} s before the last page leg finished "
                         f"({last.astimezone(dt.timezone(dt.timedelta(hours=tz_offset))):%H:%M:%S} local); "
                         f"it read a half-written feed (aged: {aged or '-'})"))
        elif t10 < first:
            out.append(F("b", "DECISION", "cockpit_predates_latest_sync", "WARN",
                         int((first - t10).total_seconds() / 60),
                         "Top 10 has not run since the latest sync; re-run it"))
        else:
            out.append(F("b", "DECISION", "cockpit_after_sync", "PASS", 0, "Top 10 ran after the sync completed"))
    out.append(F("b", "DECISION", "cockpit_output", "INFO", None, f"output: {output}; aged: {aged or '-'}"))
    # Qualified / executable
    qual = []
    for i, r in enumerate(rows):
        if _s(r[0]).startswith("ALL QUALIFIED"):
            for rr in rows[i + 2:]:
                if _num(rr[0]) is None:
                    break
                qual.append((_s(rr[1]), _s(rr[4]), _s(rr[13]) if len(rr) > 13 else ""))
            break
    sectors = collections.Counter(q[1] for q in qual)
    metrics.update(qualified=len(qual), qualified_selected=sum(1 for q in qual if q[2].lower() == "yes"),
                   qualified_sectors=dict(sectors.most_common(3)))
    if qual and sectors and sectors.most_common(1)[0][1] / len(qual) > 0.5:
        out.append(F("b", "DECISION", "qualified_set_sector_concentration", "WARN", sectors.most_common(1)[0][1],
                     f"{sectors.most_common(1)[0][0]} is {sectors.most_common(1)[0][1]} of {len(qual)} qualified plans"))
    # Caps vs policy
    want_n, want_pct = sector_cap
    t10_n = _num(_panel_value(rows, "T10: Max Per Sector"))
    prow = book.rows("Portfolio_Decision") or []
    pf_pct = _num(_panel_value(prow, "PF: Max Sector %"))
    metrics.update(t10_max_per_sector=t10_n, pf_max_sector_pct=pf_pct)
    bad = []
    if t10_n is not None and t10_n != want_n:
        bad.append(f"T10 Max Per Sector {t10_n:g} (policy {want_n})")
    if pf_pct is not None and pf_pct != want_pct:
        bad.append(f"PF Max Sector % {pf_pct:g} (policy {want_pct})")
    out.append(F("b", "DECISION", "sector_caps_vs_policy", "PASS" if not bad else "FAIL", len(bad),
                 f"policy {want_n} names / {want_pct} %", bad))
    return out, metrics


# --------------------------------------------------------------------------- #
# Gate (c): versions
# --------------------------------------------------------------------------- #
def check_versions(book, tz_offset, manifest):
    out, v = [], {}
    pages, globals_ = _status_pages(book, tz_offset)
    for p, d in pages.items():
        if d["version"]:
            v["run_dashboard_sync"] = d["version"]
    rows = book.rows("Dashboard_Audit") or []
    for r in rows[:4]:
        for c in r:
            m = re.search(r"engine=core\.data_engine_v2=([\d.]+)", _s(c))
            if m:
                v["data_engine_v2"] = m.group(1)
    st = _status_line(book.rows("Top_10_Investments") or [])
    for key, pat in (("opportunity_builder", r"builder v([\d.]+)"), ("route", r"route v([\d.]+)")):
        m = re.search(pat, st)
        if m:
            v[key] = m.group(1)
    st = _status_line(book.rows("Portfolio_Decision") or [])
    m = re.search(r"actions v([\d.]+)", st)
    if m:
        v["portfolio_actions"] = m.group(1)
    s1 = book.rows("S1_Gate") or []
    m = re.search(r"S-1 GATE v([\d.]+)", _s(s1[0][0]) if s1 else "")
    if m:
        v["run_shadow_scorer"] = m.group(1)
    cal = book.rows("_S1_Calibration") or []
    if len(cal) > 1 and len(cal[1]) > 9 and _s(cal[1][9]):
        v["track_performance"] = _s(cal[1][9])
    sb = book.rows("Shadow_Board") or []
    m = re.search(r"SHADOW BOARD v([\d.]+)", _s(sb[0][0]) if sb else "")
    if m:
        v["run_shadow_board"] = m.group(1)
    out.append(F("c", "PIPELINE", "version_stamps", "INFO", len(v), ", ".join(f"{k} {x}" for k, x in sorted(v.items()))))
    if manifest:
        bad = [f"{k}: workbook {v.get(k)} vs manifest {x}" for k, x in manifest.items() if k in v and v[k] != x]
        missing = [k for k in manifest if k not in v]
        out.append(F("c", "PIPELINE", "versions_vs_manifest", "PASS" if not bad else "FAIL", len(bad),
                     f"{len(manifest) - len(missing)} compared; not stamped in workbook: {missing}", bad))
    for k, val in globals_.items():
        if k.startswith("TFB "):
            v[f"global:{k}"] = val
    return out, v


# --------------------------------------------------------------------------- #
# Gate (d): forecast vs reality
# --------------------------------------------------------------------------- #
def check_model(book, tz_offset):
    out, metrics = [], {}
    ix, body = book.table("Performance_Log", header_pred=lambda r: _s(r[0]) == "Record ID")
    if ix:
        recs = [r for r in body if len(_s(r[0])) == 36 and _s(r[ix.get("Key", 1)])]
        st_i, oc_i, key_i = ix.get("Status"), ix.get("Outcome"), ix.get("Key", 1)
        matured = [r for r in recs if st_i is not None and _s(r[st_i]).lower() == "matured"]
        wins = sum(1 for r in matured if oc_i is not None and _s(r[oc_i]).upper() == "WIN")
        losses = sum(1 for r in matured if oc_i is not None and _s(r[oc_i]).upper() == "LOSS")
        wr = wins / (wins + losses) if wins + losses else None
        kc = collections.Counter(_s(r[key_i]) for r in matured)
        dups = sum(v - 1 for v in kc.values() if v > 1)
        risk_cols = [h for h in ix if any(k in h for k in ("Volatility", "Drawdown", "Sharpe"))]
        zero_cols = [h for h in risk_cols if recs and all((_num(r[ix[h]]) or 0) == 0 for r in recs)]
        metrics.update(perf_records=len(recs), perf_matured=len(matured), perf_wins=wins, perf_losses=losses,
                       perf_win_rate=(round(wr * 100, 2) if wr is not None else None), perf_dup_matured_keys=dups,
                       perf_zero_risk_cols=zero_cols)
        out.append(F("d", "MODEL", "performance_log_matured_win_rate",
                     "PASS" if wr is not None and wr >= 0.5 else "FAIL",
                     (round(wr * 100, 2) if wr is not None else None),
                     f"{wins} wins / {losses} losses over {len(matured)} matured records (breakeven excluded)"))
        out.append(F("d", "MODEL", "performance_log_duplicate_matured_keys", "PASS" if not dups else "FAIL", dups,
                     "P-189; dedup fix exists in track_performance v6.41.0 (TRACK_DEDUP_MATURED)"))
        out.append(F("d", "MODEL", "performance_log_risk_columns_all_zero", "PASS" if not zero_cols else "FAIL",
                     len(zero_cols), "P-188; zero is not 'unavailable'", zero_cols))
    ix, body = book.table("Signal_History")
    if ix and "Symbol" in ix:
        d_i = next((ix[h] for h in ix if h.startswith("Date")), None)
        if d_i is not None:
            c = collections.Counter((_date_of(r[d_i]), _s(r[ix["Symbol"]])) for r in body)
            multi = sum(1 for v in c.values() if v > 1)
            metrics.update(signal_rows=len(body), signal_symbol_days=len(c), signal_multi_version_days=multi,
                           signal_max_versions=max(c.values()) if c else 0)
            out.append(F("d", "MODEL", "signal_history_multi_version_symbol_days",
                         "PASS" if multi == 0 else "WARN", multi,
                         f"of {len(c)} symbol-days hold more than one version (max {max(c.values()) if c else 0}); P-191"))
    cal = book.rows("_S1_Calibration") or []
    if len(cal) > 1 and _s(cal[1][1]):
        state, n, mae, signed, band = _s(cal[1][1]), _num(cal[1][2]), _num(cal[1][3]), _num(cal[1][4]), _num(cal[1][5])
        metrics.update(calibration={"state": state, "n": n, "mae_pp": mae, "signed_pp": signed, "band_pp": band})
        out.append(F("d", "MODEL", "calibration_state", "PASS" if state.upper() == "PASS" else "FAIL", mae,
                     f"mean |err| {mae} pp (signed {signed} pp) vs band {band} pp, n={n:g}" if n else state))
    _, globals_ = _status_pages(book, tz_offset)
    m = re.search(r"brier=([\d.]+)", globals_.get("TFB Calibration", ""))
    if m:
        brier = float(m.group(1))
        metrics["brier"] = brier
        out.append(F("d", "MODEL", "brier_score", "PASS" if brier < 0.25 else "FAIL", brier,
                     "no-skill line is 0.25; lower is better"))
    s1 = book.rows("S1_Gate") or []
    if s1:
        verdict = (re.search(r"verdict: (\w+)", " ".join(_s(c) for c in s1[0])) or [None, None])[1]
        crit = {}
        for r in s1:
            if _num(r[0]) is not None and len(r) > 2 and _s(r[2]):
                crit[int(_num(r[0]))] = (_s(r[1]), _s(r[2]), _s(r[3]) if len(r) > 3 else "")
        metrics.update(s1_verdict=verdict, s1_criteria={k: v[1] for k, v in crit.items()})
        out.append(F("d", "MODEL", "s1_gate_verdict", "INFO", None,
                     f"{verdict}; " + "; ".join(f"{k} {v[0][:28]}: {v[1]}" for k, v in sorted(crit.items()))))
        ca_rows = len(book.rows("_Corporate_Actions") or [])
        ca_crit = next((v for v in crit.values() if "corporate" in v[0].lower()), None)
        if ca_crit:
            empty = ca_rows <= 1
            out.append(F("d", "GOVERNANCE", "s1_corporate_actions_criterion_on_empty_table",
                         "FAIL" if (empty and ca_crit[1].upper() == "PASS") else "PASS", ca_rows - 1 if ca_rows else 0,
                         f"criterion reads {ca_crit[1]} ('{ca_crit[2][:40]}') while _Corporate_Actions has "
                         f"{max(ca_rows - 1, 0)} event rows"))
    ix, body = book.table("Hypothesis_Registry", header_pred=lambda r: _s(r[0]) == "Hypothesis ID")
    if ix and "Status" in ix:
        sc = collections.Counter(_s(r[ix["Status"]]) for r in body if _s(r[0]))
        insuff = sc.get("INSUFFICIENT_DATA", 0)
        metrics["hypotheses"] = dict(sc)
        out.append(F("d", "MODEL", "hypotheses_without_verdict", "PASS" if insuff == 0 else "WARN", insuff,
                     f"of {sum(sc.values())} hypotheses are INSUFFICIENT_DATA (0 trigger events)"))
    return out, metrics


# --------------------------------------------------------------------------- #
# Gate (e): _Run_Log guard health
# --------------------------------------------------------------------------- #
def _cron_slots(cron):
    m = re.match(r"^\s*(\d{1,2})\s+([\d,]+)\s*$", cron)
    if not m:
        raise ValueError("cron must look like '17 4,12,20' (minute, hours UTC)")
    return int(m.group(1)), sorted(int(h) for h in m.group(2).split(",") if h)


def check_runlog(book, tz_offset, since_days, eodhd_target, cron, now_utc, asof_date):
    out, metrics = [], {}
    rows = book.rows("_Run_Log") or []
    if not rows:
        return [F("e", "PIPELINE", "run_log_present", "WARN", 0, "_Run_Log not in export")], metrics
    since = now_utc - dt.timedelta(days=since_days)
    recent, levels, errors = [], collections.Counter(), collections.Counter()
    quota = {}
    rows402 = 0
    refusals = 0
    sync_rows = []   # (ts, run id or None) for run_dashboard_sync rows in the window
    for r in rows[1:]:
        ts = _parse_dt(r[0], tz_offset)
        if ts is None:
            continue
        lvl, action, page, msg = _s(r[1]).upper(), _s(r[2]), _s(r[3]), _s(r[5]) if len(r) > 5 else ""
        details = _s(r[9]) if len(r) > 9 else ""
        if ts >= since:
            recent.append(r)
            levels[lvl] += 1
            if lvl in ("ERROR", "CRITICAL"):
                errors[(action, msg[:60])] += 1
            if "REFUS" in msg.upper() or "REFUS" in details.upper():
                refusals += 1
        m = re.search(r"\[EODHD-QUOTA[^\]]*\].*?used=(\d+)/(\d+).*?date=(\S+)", msg)
        if m:
            used, limit, date = int(m.group(1)), int(m.group(2)), m.group(3)
            quota[date] = max(quota.get(date, 0), used)
            m2 = re.search(r"rows402 new=(\d+)", msg)
            if m2 and ts >= since:
                rows402 += int(m2.group(1))
        if action == "run_dashboard_sync" and ts >= since:
            rid = (re.search(r"run=(\d+)", msg) or re.search(r'"run_id":\s*"?(\d+)', details) or [None, None])[1]
            sync_rows.append((ts, rid))
    # A run = a cluster of sync rows with no gap > 90 min; the first row of a run rarely carries the run id.
    run_first = {}
    cluster, cid = [], None
    for ts, rid in sorted(sync_rows, key=lambda x: x[0]):
        if cluster and (ts - cluster[-1]).total_seconds() > 90 * 60:
            run_first[cid or f"run@{cluster[0]:%m-%dT%H:%MZ}"] = cluster[0]
            cluster, cid = [], None
        cluster.append(ts)
        cid = cid or rid
    if cluster:
        run_first[cid or f"run@{cluster[0]:%m-%dT%H:%MZ}"] = cluster[0]
    metrics.update(run_log_rows=len(rows) - 1, run_log_recent=len(recent), levels=dict(levels))
    out.append(F("e", "PIPELINE", "run_log_errors", "PASS" if not errors else "WARN", sum(errors.values()),
                 f"ERROR/CRITICAL rows in the last {since_days} days",
                 [f"{c}x {a}: {m}" for (a, m), c in errors.most_common(5)]))
    out.append(F("e", "PIPELINE", "identity_guard_refusals_logged", "INFO", refusals,
                 "refusal rows in _Run_Log (Render-side refusals that never reach the Sheet are invisible here)"))
    # EODHD per counter date
    recent_dates = [d for d in sorted(quota) if d >= (now_utc - dt.timedelta(days=since_days)).strftime("%Y-%m-%d")]
    over = [f"{d} {quota[d]:,}" for d in recent_dates if quota[d] > eodhd_target and d != asof_date]
    metrics["eodhd_daily_max"] = {d: quota[d] for d in recent_dates}
    out.append(F("e", "PIPELINE", "eodhd_daily_over_target", "PASS" if not over else "FAIL", len(over),
                 f"of {len([d for d in recent_dates if d != asof_date])} full counter dates in the window are above "
                 f"{eodhd_target:,} calls (today excluded as partial)", over))
    out.append(F("e", "PIPELINE", "eodhd_http_402_rows", "PASS" if rows402 == 0 else "FAIL", rows402,
                 f"rows that hit HTTP 402 in the last {since_days} days"))
    # Schedule lateness
    try:
        minute, hours = _cron_slots(cron)
    except ValueError as e:
        out.append(F("e", "PIPELINE", "schedule_lateness", "WARN", None, str(e)))
        return out, metrics
    late = []
    worst = 0
    for rid, start in sorted(run_first.items(), key=lambda kv: kv[1]):
        cands = []
        for dd in (0, 1):
            day = (start - dt.timedelta(days=dd)).date()
            for h in hours:
                cands.append(dt.datetime(day.year, day.month, day.day, h, minute, tzinfo=dt.timezone.utc))
        planned = max((c for c in cands if c <= start), default=None)
        if planned is None:
            continue
        mins = int((start - planned).total_seconds() // 60)
        if mins > 12 * 60:
            continue  # not attributable to a slot (manual dispatch)
        worst = max(worst, mins)
        late.append(f"run {rid}: slot {planned:%m-%d %H:%M}Z started {start:%H:%M}Z, +{mins // 60}h{mins % 60:02d}m")
    metrics["schedule_worst_lateness_min"] = worst
    out.append(F("e", "PIPELINE", "schedule_lateness", "PASS" if worst <= 15 else ("WARN" if worst <= 60 else "FAIL"),
                 worst, f"worst start delay vs cron '{cron}' over {len(late)} runs (first _Run_Log row of each run); P-170",
                 late[-6:]))
    da = book.rows("Dashboard_Audit") or []
    capped = [f"{_s(r[0])} {_num(r[3]):g}" for r in da if _s(r[1]) == "scope.coverage" and _s(r[2]) == "WARN"]
    if da:
        out.append(F("e", "PIPELINE", "dashboard_audit_scope_capped", "PASS" if not capped else "WARN", len(capped),
                     "in-workbook validator samples a capped row count (P-190)", capped))
    return out, metrics


# --------------------------------------------------------------------------- #
# Gate (f): verdict
# --------------------------------------------------------------------------- #
LANES = ["DATA", "BOOK", "DECISION", "PIPELINE", "MODEL", "GOVERNANCE"]
RANK = {"PASS": 0, "INFO": 0, "NOT_CHECKED": 0, "WARN": 1, "FAIL": 2}


def verdicts(findings):
    lanes = {}
    for lane in LANES:
        fs = [f for f in findings if f["lane"] == lane]
        worst = max((RANK[f["status"]] for f in fs), default=0)
        lanes[lane] = {"rag": ["GREEN", "AMBER", "RED"][worst],
                       "fail": sum(1 for f in fs if f["status"] == "FAIL"),
                       "warn": sum(1 for f in fs if f["status"] == "WARN"),
                       "checks": len(fs)}
    overall = ["GREEN", "AMBER", "RED"][max(RANK[f["status"]] for f in findings)] if findings else "GREEN"
    return lanes, overall


def audit_workbook(path, expect=None, asof=None, dq=80.0, rel=70.0, sector_cap=(3, 40.0),
                   eodhd_target=90000, cron="17 4,12,20", tz_offset=3, since_days=7,
                   manifest=None, now_utc=None):
    """Run every gate on one export. Returns a JSON-serialisable result dict."""
    expect = dict(DEFAULT_EXPECT if expect is None else expect)
    now_utc = now_utc or dt.datetime.now(dt.timezone.utc)
    book = Book(path)
    # as-of date = the most common Last Updated (UTC) date across market pages unless given
    if not asof:
        dates = collections.Counter()
        for page in MARKET_PAGES:
            ix, body = book.table(page)
            if ix and "Last Updated (UTC)" in ix:
                dates.update(_date_of(r[ix["Last Updated (UTC)"]]) for r in body)
        dates.pop(None, None)
        asof = max(dates) if dates else now_utc.strftime("%Y-%m-%d")
    findings, metrics = [], {"asof": asof}
    findings += check_structure(book, now_utc)
    f, m = check_market_pages(book, expect, asof, dq, rel)
    findings += f
    metrics["pages"] = m
    findings += check_crypto_identity(book)
    findings += check_forecast_pairs(book)
    findings += check_horizon(book)
    findings += check_pit(book)
    f, m = check_portfolio(book, tz_offset, since_days, now_utc)
    findings += f
    metrics["book"] = m
    f, m = check_cockpit(book, tz_offset, sector_cap)
    findings += f
    metrics["cockpit"] = m
    f, m = check_versions(book, tz_offset, manifest)
    findings += f
    metrics["versions"] = m
    f, m = check_model(book, tz_offset)
    findings += f
    metrics["model"] = m
    f, m = check_runlog(book, tz_offset, since_days, eodhd_target, cron, now_utc, asof)
    findings += f
    metrics["runlog"] = m
    findings.append(F("gov", "GOVERNANCE", "workbook_sharing_acl", "NOT_CHECKED", None,
                      "cannot be read from an export; check Drive sharing by hand (A1/C01)"))
    lanes, overall = verdicts(findings)
    fixes = [f"{f['lane']}: {f['check']} ({f['count']}) - {f['detail']}" for f in findings if f["status"] == "FAIL"]
    return {"script": "tfb_export_audit", "version": SCRIPT_VERSION, "export": os.path.basename(path),
            "export_sha256": _sha256_file(path), "asof": asof, "overall": overall, "lanes": lanes,
            "findings": findings, "metrics": metrics, "fix_list": fixes,
            "params": {"expect": expect, "dq": dq, "rel": rel, "sector_cap": list(sector_cap),
                       "eodhd_target": eodhd_target, "cron": cron, "tz_offset": tz_offset,
                       "since_days": since_days, "manifest": bool(manifest)}}


def _jsonable(o):
    """Stringify dict keys (None / numbers) so sort_keys and json.dump never trip."""
    if isinstance(o, dict):
        return {str(k): _jsonable(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [_jsonable(v) for v in o]
    return o


def digest(result):
    """sha256 over the result minus anything time-of-run dependent."""
    body = _jsonable({k: v for k, v in result.items() if k not in ("generated_at", "digest")})
    return hashlib.sha256(json.dumps(body, sort_keys=True, default=str).encode()).hexdigest()[:16]


def render_md(res):
    L = [f"# TFB export audit v{SCRIPT_VERSION} - {res['export']}",
         "", f"As of {res['asof']} - export sha256 `{res['export_sha256'][:16]}...` - overall **{res['overall']}**", "",
         "| Lane | RAG | FAIL | WARN | Checks |", "|---|---|---|---|---|"]
    for lane in LANES:
        v = res["lanes"][lane]
        L.append(f"| {lane} | {v['rag']} | {v['fail']} | {v['warn']} | {v['checks']} |")
    L += ["", "## Fix list (FAIL)", ""]
    L += [f"{i + 1}. {x}" for i, x in enumerate(res["fix_list"])] or ["(none)"]
    gates = [("a", "Gate (a) - data integrity"), ("b", "Gate (b) - recommendation coherence"),
             ("c", "Gate (c) - cross-version consistency"), ("d", "Gate (d) - forecast vs reality"),
             ("e", "Gate (e) - guard health (_Run_Log)"), ("book", "Book - ledger and cash"),
             ("gov", "Governance")]
    for g, title in gates:
        fs = [f for f in res["findings"] if f["gate"] == g]
        if not fs:
            continue
        L += ["", f"## {title}", "", "| Check | Status | Count | Detail | Examples |", "|---|---|---|---|---|"]
        for f in fs:
            ex = "; ".join(str(e) for e in f["examples"][:4])
            L.append(f"| {f['check']} | {f['status']} | {'' if f['count'] is None else f['count']} | "
                     f"{f['detail'].replace('|', '/')} | {ex.replace('|', '/')} |")
    pages = res["metrics"].get("pages", {})
    if pages:
        L += ["", "## Page metrics", "", "| Page | Rows | Fresh | No price | No name | PM > 1 | INVEST / eligible |",
              "|---|---|---|---|---|---|---|"]
        for p in MARKET_PAGES:
            m = pages.get(p)
            if m:
                L.append(f"| {p} | {m['rows']} | {m.get('fresh_rows')} | {m['invalid_price']} | {m['missing_names']} | "
                         f"{m['pm_over_1']}/{m['pm_n']} | {m['invest']} / {m['invest_eligible']} |")
    b = res["metrics"].get("book", {})
    if b:
        L += ["", "## Book", "",
              f"- Active lots {b.get('active_lots')}; holdings {b.get('holdings_sar'):,} SAR; cash {b.get('cash_sar')} SAR; "
              f"NAV {b.get('nav_sar')} SAR; ledger totals {b.get('ledger_totals')}",
              f"- Portfolio_Decision KPI: {b.get('pd_kpi')}",
              f"- Fills in window {b.get('fills_in_window')}; trade notes {b.get('trade_notes')}"]
    v = res["metrics"].get("versions", {})
    if v:
        L += ["", "## Versions", ""] + [f"- {k}: {x}" for k, x in sorted(v.items())]
    return "\n".join(L) + "\n"


# --------------------------------------------------------------------------- #
# Self-test: synthetic workbooks through the REAL audit path
# --------------------------------------------------------------------------- #
def _market_header():
    base = ["Symbol", "Name", "Asset Class", "Exchange", "Currency", "Country", "Sector", "Industry",
            "Current Price", "Previous Close", "Open", "Day High", "Day Low", "52W High", "52W Low",
            "Gross Margin", "Operating Margin", "Profit Margin", "Forecast Price 1M", "Forecast Price 3M",
            "Forecast Price 12M", "Expected ROI 12M", "Forecast Confidence", "Recommendation", "Final Action",
            "Data Quality Score", "Forecast Reliability Score", "Horizon Days", "Invest Period Label",
            "Last Updated (UTC)"]
    return base + [f"Col{i}" for i in range(len(base), MAIN_WIDTH)]


def _row(h, **kw):
    r = [None] * len(h)
    for k, v in kw.items():
        r[h.index(k)] = v
    return r


def _build_synthetic(path, clean, asof="2026-10-03", tz_offset=3):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    H = _market_header()
    upd = f"{asof} 10:00:00"
    old = "2026-09-20 10:00:00"
    ok = dict(Name="Good Co", **{"Current Price": 10.0, "Previous Close": 9.9, "52W High": 12.0, "52W Low": 8.0,
                                 "Profit Margin": 0.12, "Forecast Price 12M": 11.0, "Expected ROI 12M": 0.1,
                                 "Forecast Confidence": 0.7, "Recommendation": "BUY", "Final Action": "INVEST",
                                 "Data Quality Score": 90, "Forecast Reliability Score": 80, "Horizon Days": 90,
                                 "Invest Period Label": "3M", "Last Updated (UTC)": upd})
    def page(name, rows):
        ws = wb.create_sheet(name)
        ws.append(H)
        for r in rows:
            ws.append(r)
    ml = [_row(H, Symbol=f"A{i}.SR", **ok) for i in range(5)]
    gm = [_row(H, Symbol=f"G{i}.US", **ok) for i in range(10)]
    cfx = [_row(H, Symbol="BTC-USD", **{**ok, "Name": "Bitcoin USD", "Final Action": "HOLD"}),
           _row(H, Symbol="GC=F", **{**ok, "Name": "Gold", "Final Action": "HOLD"})]
    mf = [_row(H, Symbol=f"F{i}.US", **{**ok, "Final Action": "HOLD"}) for i in range(4)]
    if not clean:
        gm[0][H.index("Symbol")] = "G1.US"                       # duplicate symbol
        gm[2][H.index("Name")] = None                            # missing name
        gm[3][H.index("Current Price")] = 0                      # invalid price
        gm[4][H.index("Previous Close")] = 5.0                   # +100 % move
        gm[5][H.index("Current Price")] = 13.0                   # outside 52W
        gm[5][H.index("Previous Close")] = 12.9
        gm[5][H.index("Forecast Price 12M")] = 14.3               # keep its forecast/ROI pair consistent
        gm[6][H.index("Profit Margin")] = 25.8                   # unit error
        gm[7][H.index("Expected ROI 12M")] = -0.10                # pair mismatch (f12 11 @ 10 = +10 %)
        gm[8][H.index("Data Quality Score")] = 50                 # INVEST not eligible
        gm[9][H.index("Recommendation")] = "SELL"                 # SELL-class INVEST
        gm[9][H.index("Horizon Days")] = 365                      # horizon mismatch
        gm[1][H.index("Last Updated (UTC)")] = old                # stale row
        cfx.append(_row(H, Symbol="SUI-USD", **{**ok, "Name": "Salmonation USD", "Final Action": "HOLD"}))
        cfx.append(_row(H, Symbol="SHIB-USD", **{**ok, "Name": "Shiba Inu USD", "Final Action": "HOLD",
                                               "Current Price": 0.00001, "Previous Close": 0.00001,
                                               "52W High": 0.00002, "52W Low": 0.000005,
                                               "Forecast Price 1M": 0.0, "Forecast Price 12M": 0.0,
                                               "Expected ROI 12M": None}))
    page("Market_Leaders", ml); page("Global_Markets", gm); page("Commodities_FX", cfx); page("Mutual_Funds", mf)
    expect = {"Market_Leaders": len(ml), "Global_Markets": len(gm), "Commodities_FX": len(cfx), "Mutual_Funds": len(mf)}
    # My_Portfolio (122 cols)
    PH = ["Symbol", "Name", "Asset Class", "Exchange", "Currency", "Country", "Sector", "Industry", "Current Price",
          "Horizon Days", "Invest Period Label"] + [f"P{i}" for i in range(11, PORTFOLIO_WIDTH)]
    ws = wb.create_sheet("My_Portfolio"); ws.append(PH)
    for sym, px in (("AAA.US", 10.0), ("BBB.US", 20.0)):
        r = [None] * len(PH); r[0] = sym; r[4] = "USD"; r[8] = px
        r[9] = 90 if clean else 365; r[10] = "3M"
        ws.append(r)
    # Ledger
    ws = wb.create_sheet("_Portfolio_CostBasis")
    ws.append(["_Portfolio_CostBasis - INVESTMENT LEDGER"])
    ws.append(["Status:", "Last refresh 2026-10-03 13:39:00 | 2 active · 1 closed | totals: closed -100 · active 50 · lifetime -50 SAR"])
    ws.append(["legend"])
    ws.append(["Symbol", "Name", "Ccy", "Status", "Buy Date", "Buy Price", "Shares", "Buy Fees", "Sell Date",
               "Sell Price", "Sell Fees", "Dividends Recv", "Notes", "Cost Basis"])
    ws.append(["AAA.US", "A", "USD", "Active", dt.datetime(2026, 9, 28), 9.0, 100, 2.29, None, None, None, 0, None, 900])
    ws.append(["BBB.US", "B", "USD", "Active", dt.datetime(2026, 9, 1), 19.0, 10, 2.29, None, None, None, 0, None, 190])
    ws.append(["YUM", "Y", "USD", "Inactive", dt.datetime(2026, 8, 12), 144.71, 24, None, dt.datetime(2026, 10, 1),
               136.01, 2.29, 18, None, 3473.04])
    # Cash
    ws = wb.create_sheet("_Cash_Snapshot")
    ws.append(["Date", "Time", "Type", "Amount SAR", "Balance SAR", "Delta vs Prev", "Note"])
    ws.append([dt.datetime(2026, 10, 1), None, "SNAPSHOT", None, 34166.25, None, "USD 9130"])
    ws.append([dt.datetime(2026, 10, 1 if not clean else 2), None, "SNAPSHOT", None, 46398.75, None, "USD 9130" if not clean else "USD 12373"])
    cash = 46398.75
    holdings = 100 * 10.0 * 3.75 + 10 * 20.0 * 3.75   # 4,500
    # Portfolio_Decision
    ws = wb.create_sheet("Portfolio_Decision")
    ws.append(["MY PORTFOLIO - DECISION"])
    ws.append(["Status:", "Last run 2026-10-03 13:10:36 | status: ok | route v4.16.0 | actions v1.14.0 | builder v1.23.0"])
    ws.append([])
    ws.append(["CONTROL PANEL"])
    ws.append(["PF: Cash Available SAR", cash, None, "PF: Target Cash %", 10, None, "PF: Max Position %", 20])
    ws.append(["PF: Max Sector %", 40 if clean else 30])
    ws.append([])
    ws.append(["KPIs"])
    ws.append(["Portfolio (SAR)", "Holdings (SAR)", "Cash (SAR)", "Cash %"])
    ws.append([holdings + cash if clean else holdings + 34166.25, holdings, cash if clean else 34166.25, 50])
    ws.append([])
    ws.append(["ACTIONS"])
    ws.append(["Action", "Symbol", "Name", "Sector", "Ccy", "Qty"])
    ws.append(["HOLD", "AAA.US", "A", "X", "USD", 100])
    ws.append(["HOLD", "BBB.US", "B", "X", "USD", 10])
    if not clean:
        ws.append(["HOLD", "YUM", "Y", "X", "USD", 24])
    ws.append([None])
    # Top 10
    gm_done = f"{asof} 13:48:28+03:00"
    t10_run = f"{asof} 13:47:12" if not clean else f"{asof} 13:50:00"
    ws = wb.create_sheet("Top_10_Investments")
    ws.append(["TOP 10 INVESTMENTS"])
    ws.append(["Status:", f"Last run {t10_run} | status: ok | output: {'WITHHELD' if not clean else 'OK'} | route v4.16.0 | builder v1.23.0"
                          + ("" if clean else " | aged:GM")])
    ws.append(["CONTROL PANEL"])
    ws.append(["T10: Max Per Sector", 3 if clean else 2, "T10: Max Per Market", 10])
    ws.append(["ALL QUALIFIED - INVEST opportunity set (3)"])
    ws.append(["Rank", "Symbol", "Name", "Market", "Sector", "ROI", "E", "A", "R", "Rel", "DQ", "Risk", "Score", "Selected", "Why"])
    ws.append([1, "H1.US", "n", "m", "Financials", 1, 1, 1, 1, 70, 100, "Low", 70, "Yes", "-"])
    ws.append([2, "H2.US", "n", "m", "Financials", 1, 1, 1, 1, 70, 100, "Low", 70, "No", "-"])
    ws.append([3, "H3.US", "n", "m", "Energy", 1, 1, 1, 1, 70, 100, "Low", 70, "No", "-"])
    # _Status
    ws = wb.create_sheet("_Status")
    ws.append(["Page", "Last Updated", "Status", "Message", "Endpoint", "HTTP", "Rows", "Cols", "ms", "W", None, "Global Key", "Value"])
    ws.append(["Market_Leaders", f"{asof} 12:59:18+03:00", "SUCCESS", "[STATUS-STAMP v6.64.0] leg=success run=37114669145",
               None, None, None, None, None, None, None, "TFB Calibration", "BLOCKED:0.458 | brier=0.2811 | as_of=x" if not clean else "brier=0.2400"])
    ws.append(["Global_Markets", gm_done, "SUCCESS", "[STATUS-STAMP v6.64.0] leg=success run=37114669145"])
    ws.append(["Commodities_FX", f"{asof} 13:01:47+03:00", "SUCCESS", "[STATUS-STAMP v6.64.0] run=37114669145"])
    ws.append(["Mutual_Funds", f"{asof} 13:20:14+03:00", "SUCCESS", "[STATUS-STAMP v6.64.0] run=37114669145"])
    # S1 gate + CA + calibration + hypotheses + PIT
    ws = wb.create_sheet("S1_Gate")
    ws.append(["S-1 GATE v1.9.1", "as of 2026-10-02", "verdict: NOT_DECIDABLE"])
    ws.append(["#", "Criterion", "Status", "Detail"])
    ws.append([1, "4+ weeks shadow evidence", "PENDING", "14/28"])
    ws.append([5, "corporate-actions + point-in-time", "PASS" if not clean else "NOT_EVALUABLE", "CA clean"])
    ws = wb.create_sheet("_Corporate_Actions"); ws.append(["Symbol", "Type", "Date"])
    ws = wb.create_sheet("_S1_Calibration")
    ws.append(["As Of", "State", "N", "MAE", "Signed", "Band", "Min", "By", "Detail", "Writer Version"])
    ws.append(["2026-10-03", "PASS", 6047, 3.13, -0.39, 10, 20, "", "", "6.41.0"])
    ws = wb.create_sheet("Hypothesis_Registry")
    for _ in range(4):
        ws.append([])
    ws.append(["Hypothesis ID", "Name", "Type", "Source", "Trigger", "Dir", "H", "Min", "E", "T", "Status"])
    ws.append(["H1", "n", "t", "s", "v", "d", 20, 30, 50, 2, "INSUFFICIENT_DATA" if not clean else "SUPPORTED"])
    ws = wb.create_sheet("_PIT_Fundamentals")
    ws.append(["Symbol", "As Of", "Forecast Reliability"])
    ws.append(["AAA.US", "2026-07-21", "2026-07-21 21:23:10" if not clean else 71.5])
    ws.append(["BBB.US", "2026-07-21", 70.0])
    # Performance_Log + Signal_History
    ws = wb.create_sheet("Performance_Log")
    for _ in range(4):
        ws.append([])
    ws.append(["Record ID", "Key", "Symbol", "Horizon", "Date", "Entry", "Rec", "Score", "Risk", "Conf", "Origin",
               "Target", "ROI", "TDate", "Status", "Cur", "Unreal", "Real", "Outcome", "Volatility", "Max Drawdown", "Sharpe Ratio"])
    def rec(i, key, status, outcome, vol, dd=0, sharpe=0):
        return [f"{i:08d}-0000-0000-0000-000000000000", key, key.split("|")[0], "1W", "2026-09-01", 1, "HOLD", 50, "LOW",
                "HIGH", "x", 1, 0, "2026-09-08", status, 1, 0, 0, outcome, vol, dd, sharpe]
    if clean:
        ws.append(rec(1, "A|1W|1", "matured", "WIN", 0.2, -0.05, 1.1)); ws.append(rec(2, "B|1W|1", "matured", "WIN", 0.3, -0.08, 0.9))
        ws.append(rec(3, "C|1W|1", "matured", "LOSS", 0.1, -0.02, -0.4))
    else:
        ws.append(rec(1, "A|1W|1", "matured", "WIN", 0)); ws.append(rec(2, "A|1W|1", "matured", "LOSS", 0))
        ws.append(rec(3, "B|1W|1", "matured", "LOSS", 0)); ws.append(rec(4, "C|1W|1", "active", "", 0))
    ws = wb.create_sheet("Signal_History")
    ws.append(["Snapshot ID", "Key", "Symbol", "Date (Riyadh)"])
    ws.append(["s1", "A|20261001", "A", "2026-10-01"])
    ws.append(["s2", "A|20261001", "A", "2026-10-01" if not clean else "2026-10-02"])
    # Run log
    ws = wb.create_sheet("_Run_Log")
    ws.append(["Timestamp", "Level", "Action", "Page", "Status", "Message", "Endpoint", "HTTP", "ms", "Details JSON"])
    start = dt.datetime(2026, 10, 3, 7, 20) if clean else dt.datetime(2026, 10, 3, 12, 55)   # local; slot 04:17Z = 07:17 local
    ws.append([start, "INFO", "run_dashboard_sync", "Market_Leaders", "OK",
               "[EODHD-QUOTA v6.64.0] Market_Leaders | used=1000/400000 (0.3%) date=2026-10-02 extra=0 | rows402 new=0 | state=OK"])
    ws.append([start + dt.timedelta(minutes=1), "INFO", "run_dashboard_sync", "Global_Markets", "OK",
               f"[EODHD-QUOTA v6.64.0] Global_Markets | used={'118631' if not clean else '48000'}/400000 date=2026-10-01 extra=0 | rows402 new={'0' if clean else '3'} | state=OK",
               None, None, None, '{"run_id": "37114669145"}'])
    ws.append([start + dt.timedelta(minutes=2), "INFO", "run_dashboard_sync", "Global_Markets", "OK",
               f"[EODHD-QUOTA v6.64.0] Global_Markets | used=48351/400000 date={asof} extra=0 | rows402 new=0 | state=OK",
               None, None, None, '{"run_id": "37114669145"}'])
    if not clean:
        ws.append([dt.datetime(2026, 9, 30, 16, 5), "ERROR", "tfbDispatchDailySync", "ALL", "ERROR", "dispatch failed: no_token"])
    ws = wb.create_sheet("Dashboard_Audit")
    ws.append(["Dashboard Validation", "Generated: x"]); ws.append(["registry=2.15.0", "engine=core.data_engine_v2=5.151.0"])
    ws.append(["Page", "Check", "Status", "Count", "Examples"])
    ws.append(["Global_Markets", "scope.coverage", "WARN" if not clean else "PASS", 1500 if not clean else 0, ""])
    ws = wb.create_sheet("Shadow_Board"); ws.append(["SHADOW BOARD v1.5.0", "as of x", "equity=130,000 SAR"])
    ws = wb.create_sheet("_Trade_Notes"); ws.append(["Note ID", "Logged At", "Type", "Trade Date", "Symbol"])
    if clean:   # one note per fill in the window: AAA buy 09-28, YUM sell 10-01
        ws.append(["N1", "2026-09-28", "ENTRY", "2026-09-28", "AAA.US"])
        ws.append(["N2", "2026-10-01", "EXIT", "2026-10-01", "YUM"])
    ws = wb.create_sheet("_Decision_Diagnostics")
    ws.append(["DECISION CHAIN DIAGNOSTICS"]); ws.append([f"dd v1.4.1 | {'2026-08-11' if not clean else asof} 11:19:34 | 33 checks"])
    if not clean:
        ws = wb.create_sheet("Sheet97"); ws.append(["x"])
        ws = wb.create_sheet("_Data_Audit"); ws.append(["DATA AUDIT"]); ws.append(["#ERROR!"])
    wb.save(path)
    return expect


def selftest():
    now = dt.datetime(2026, 10, 3, 11, 0, tzinfo=dt.timezone.utc)
    checks = []
    def expect_(cond, label):
        checks.append((bool(cond), label))
    with tempfile.TemporaryDirectory() as td:
        for clean in (False, True):
            path = os.path.join(td, f"synthetic_{'clean' if clean else 'defects'}.xlsx")
            exp = _build_synthetic(path, clean)
            res = audit_workbook(path, expect=exp, asof="2026-10-03", now_utc=now, since_days=7)
            fs = {f["check"]: f for f in res["findings"]}
            def st(name):
                return fs[name]["status"] if name in fs else "MISSING"
            def cnt(name):
                return fs[name]["count"] if name in fs else None
            tag = "clean" if clean else "defects"
            if not clean:
                expect_(cnt("Global_Markets.duplicate_symbols") == 1, f"{tag}: duplicate symbol detected")
                expect_(cnt("Global_Markets.missing_names") == 1, f"{tag}: missing name")
                expect_(cnt("Global_Markets.invalid_price") == 1, f"{tag}: invalid price")
                expect_(cnt("Global_Markets.day_move_over_25pct") == 1, f"{tag}: big move")
                expect_(cnt("Global_Markets.price_outside_52w") == 1, f"{tag}: outside 52W")
                expect_(cnt("Global_Markets.profit_margin_unit_errors") == 1, f"{tag}: margin unit error")
                expect_(cnt("Global_Markets.stale_over_5_days") == 1, f"{tag}: stale row")
                expect_(cnt("Global_Markets.forecast_roi_pair_mismatch") == 1 and st("forecast_roi_pairs_total") == "FAIL",
                        f"{tag}: forecast/ROI pair mismatch")
                expect_(cnt("invest_rows_vs_eligibility") == 1 and st("invest_rows_vs_eligibility") == "FAIL"
                        and "15 rows say INVEST; 14 meet" in fs["invest_rows_vs_eligibility"]["detail"],
                        f"{tag}: INVEST vs eligibility (15 INVEST, 14 eligible, the DQ-50 row fails)")
                expect_(cnt("Global_Markets.sell_class_rows_marked_invest") == 1, f"{tag}: SELL-class INVEST")
                expect_(cnt("Global_Markets.horizon_vs_label") == 1 and cnt("My_Portfolio.horizon_vs_label") == 2,
                        f"{tag}: horizon mismatch (GM 1, portfolio 2)")
                expect_(cnt("crypto_wrong_instrument_name") == 1 and "SUI-USD" in fs["crypto_wrong_instrument_name"]["examples"][0],
                        f"{tag}: crypto wrong instrument")
                expect_(cnt("crypto_zero_forecast_price") == 1, f"{tag}: crypto zero forecast")
                expect_(cnt("pit_timestamp_in_reliability") == 1, f"{tag}: PIT timestamp")
                expect_(cnt("portfolio_decision_sold_holdings_shown") == 1 and st("portfolio_decision_sold_holdings_shown") == "FAIL",
                        f"{tag}: sold holding shown")
                expect_(st("portfolio_decision_kpi_cash_vs_snapshot") == "FAIL" and abs(cnt("portfolio_decision_kpi_cash_vs_snapshot") + 12232.5) < 0.01,
                        f"{tag}: KPI cash stale by -12,232.50")
                expect_(cnt("cash_snapshot_duplicate_dates") == 1 and cnt("cash_snapshot_stale_notes") == 1, f"{tag}: cash dup date + stale note")
                expect_(st("cockpit_ran_inside_sync_run") == "FAIL" and cnt("cockpit_ran_inside_sync_run") == 76, f"{tag}: cockpit race = 76 s")
                expect_(cnt("sector_caps_vs_policy") == 2, f"{tag}: both caps off policy")
                expect_(st("performance_log_matured_win_rate") == "FAIL" and abs(cnt("performance_log_matured_win_rate") - 33.33) < 0.01,
                        f"{tag}: win rate 33.33")
                expect_(cnt("performance_log_duplicate_matured_keys") == 1, f"{tag}: dup matured key")
                expect_(cnt("performance_log_risk_columns_all_zero") == 3, f"{tag}: 3 zero risk columns")
                expect_(cnt("signal_history_multi_version_symbol_days") == 1, f"{tag}: multi-version day")
                expect_(st("brier_score") == "FAIL" and cnt("brier_score") == 0.2811, f"{tag}: Brier 0.2811 FAIL")
                expect_(st("s1_corporate_actions_criterion_on_empty_table") == "FAIL", f"{tag}: CA criterion on empty table")
                expect_(cnt("hypotheses_without_verdict") == 1, f"{tag}: hypothesis without verdict")
                expect_(cnt("run_log_errors") == 1, f"{tag}: run-log error row")
                expect_(st("eodhd_daily_over_target") == "FAIL" and cnt("eodhd_daily_over_target") == 1, f"{tag}: EODHD over target once")
                expect_(cnt("eodhd_http_402_rows") == 3, f"{tag}: 402 rows summed")
                expect_(st("schedule_lateness") == "FAIL" and cnt("schedule_lateness") == 338, f"{tag}: lateness 5h38m = 338 min")
                expect_(cnt("dashboard_audit_scope_capped") == 1, f"{tag}: validator cap")
                expect_(st("trade_notes_vs_fills") == "FAIL" and cnt("trade_notes_vs_fills") == 2, f"{tag}: 2 fills, 0 notes")
                expect_(cnt("dead_tabs_present") == 2 and cnt("data_audit_error_cells") == 1, f"{tag}: dead tabs + #ERROR!")
                expect_(st("workbook_sharing_acl") == "NOT_CHECKED", f"{tag}: sharing reported NOT_CHECKED, never PASS")
                expect_(res["overall"] == "RED" and res["lanes"]["DATA"]["rag"] == "RED", f"{tag}: overall RED")
            else:
                fails = [f["check"] for f in res["findings"] if f["status"] == "FAIL"]
                expect_(not fails, f"{tag}: zero FAIL on a clean workbook (got {fails})")
                expect_(st("cockpit_after_sync") == "PASS", f"{tag}: cockpit after sync")
                expect_(st("portfolio_decision_kpi_cash_vs_snapshot") == "PASS", f"{tag}: KPI cash matches")
                expect_(st("schedule_lateness") == "PASS" and cnt("schedule_lateness") == 3, f"{tag}: on-time run (+3 min)")
                expect_(st("brier_score") == "PASS", f"{tag}: Brier 0.24 PASS")
                expect_(st("trade_notes_vs_fills") == "PASS", f"{tag}: notes cover fills")
                expect_(res["lanes"]["BOOK"]["rag"] == "GREEN" and res["lanes"]["MODEL"]["rag"] == "GREEN", f"{tag}: BOOK and MODEL green")
            # determinism: same workbook twice -> same digest
            res2 = audit_workbook(path, expect=exp, asof="2026-10-03", now_utc=now, since_days=7)
            expect_(digest(res) == digest(res2), f"{tag}: deterministic digest")
            # markdown renders without error and names the overall verdict
            md = render_md(res)
            expect_(res["overall"] in md and "## Fix list" in md, f"{tag}: markdown renders")
    passed = sum(1 for ok, _ in checks if ok)
    for ok, label in checks:
        print(("PASS " if ok else "FAIL ") + label)
    h = hashlib.sha256("\n".join(label for _, label in checks).encode()).hexdigest()[:16]
    print(f"[SELFTEST v{SCRIPT_VERSION}] {passed}/{len(checks)} PASS  cases-digest={h}")
    return passed == len(checks)


# --------------------------------------------------------------------------- #
def parse_args(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("export", nargs="?", help="path to the .xlsx export of the workbook")
    ap.add_argument("--json", help="write the full result JSON here")
    ap.add_argument("--md", help="write the Markdown report here")
    ap.add_argument("--expect", default=",".join(f"{k}={v}" for k, v in DEFAULT_EXPECT.items()),
                    help="expected row counts, Page=N,...")
    ap.add_argument("--asof", help="as-of date YYYY-MM-DD (default: most common Last Updated date)")
    ap.add_argument("--dq", type=float, default=80.0, help="cockpit Data Quality floor")
    ap.add_argument("--rel", type=float, default=70.0, help="cockpit reliability floor")
    ap.add_argument("--sector-cap", default="3/40", help="policy sector cap names/percent (default 3/40)")
    ap.add_argument("--eodhd-target", type=int, default=90000)
    ap.add_argument("--cron", default="17 4,12,20", help="daily_sync cron 'minute hours' in UTC")
    ap.add_argument("--tz-offset", type=int, default=3, help="workbook local offset (Riyadh = 3)")
    ap.add_argument("--since-days", type=int, default=7)
    ap.add_argument("--manifest", help="JSON {name: version} to compare stamps against")
    ap.add_argument("--now", help="override 'now' as ISO UTC (for replays)")
    ap.add_argument("--selftest", action="store_true")
    ap.add_argument("--quiet", action="store_true")
    return ap.parse_args(argv)


def main(argv=None):
    a = parse_args(argv)
    if a.selftest:
        return 0 if selftest() else 1
    if not a.export:
        sys.stderr.write("usage: tfb_export_audit.py export.xlsx [--json ...] [--md ...] | --selftest\n")
        return 2
    if not os.path.exists(a.export):
        sys.stderr.write(f"not found: {a.export}\n")
        return 2
    expect = {}
    for part in a.expect.split(","):
        if "=" in part:
            k, v = part.split("=", 1)
            expect[k.strip()] = int(v)
    n, pct = a.sector_cap.split("/")
    manifest = json.load(open(a.manifest)) if a.manifest else None
    now = _parse_dt(a.now, 0) if a.now else None
    res = audit_workbook(a.export, expect=expect, asof=a.asof, dq=a.dq, rel=a.rel,
                         sector_cap=(int(n), float(pct)), eodhd_target=a.eodhd_target, cron=a.cron,
                         tz_offset=a.tz_offset, since_days=a.since_days, manifest=manifest, now_utc=now)
    res["generated_at"] = dt.datetime.now(dt.timezone.utc).isoformat()
    res["digest"] = digest(res)
    if a.json:
        with open(a.json, "w", encoding="utf-8") as fh:
            json.dump(_jsonable(res), fh, indent=1, default=str)
    md = render_md(res)
    if a.md:
        with open(a.md, "w", encoding="utf-8") as fh:
            fh.write(md)
    if not a.quiet:
        print(md)
    lanes = " ".join(f"{k}={v['rag']}" for k, v in res["lanes"].items())
    print(f"[EXPORT-AUDIT v{SCRIPT_VERSION}] {res['export']} asof={res['asof']} overall={res['overall']} "
          f"{lanes} fails={len(res['fix_list'])} digest={res['digest']}")
    return 1 if res["fix_list"] else 0


if __name__ == "__main__":
    sys.exit(main())
