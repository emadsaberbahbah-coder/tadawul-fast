"""TFB-04 regressions through the real exported-workbook cash consumer.

This tests offline certification only. It does not authenticate broker reads,
change orders, or imply that legacy naked-cash funding inputs are verified.
"""
from __future__ import annotations

import dataclasses
import datetime as dt
from decimal import Decimal
import json
import subprocess
import sys

import openpyxl
import pytest

from scripts import tfb_export_audit as audit


NOW = dt.datetime(2026, 10, 8, 12, tzinfo=dt.timezone.utc)
HEADERS = ["Date", "Time", "Type", "Balance SAR", "Note", "Account", "Currency",
           "Balance Type", "Balance", "FX to SAR", "FX As Of"]


def _record(**changes):
    row = {"Date": dt.date(2026, 10, 8), "Time": dt.time(14, 30),
           "Type": "SNAPSHOT", "Balance SAR": 46_444.46, "Note": "broker read",
           "Account": "synthetic-account", "Currency": "USD",
           "Balance Type": "settled_cash", "Balance": 12_373.20,
           "FX to SAR": 3.753634, "FX As Of": "2026-10-08T11:25:00Z"}
    row.update(changes)
    return row


def _write(tmp_path, *records, extra_headers=()):
    path = tmp_path / "cash.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    cash = wb.create_sheet("_Cash_Snapshot")
    headers = HEADERS + list(extra_headers)
    cash.append(headers)
    for row in records:
        cash.append([row.get(header) for header in headers])
    decision = wb.create_sheet("Portfolio_Decision")
    recorded = records[-1].get("Balance SAR")
    decision.append(["PF: Cash Available SAR", recorded])
    decision.append(["Portfolio (SAR)", "Cash (SAR)"])
    decision.append([recorded, recorded])
    wb.save(path)
    return path


def _check(path, **kwargs):
    findings, metrics = audit.check_portfolio(audit.Book(path), 3, 7, NOW, **kwargs)
    return {finding["check"]: finding for finding in findings}, metrics


def _assert_uncertified(findings, metrics, reason):
    assert findings["cash_snapshot_certification"]["status"] == "FAIL"
    assert reason in metrics["cash_certification_errors"]
    assert metrics["cash_certified"] is False
    assert metrics["cash_sar"] is None
    assert metrics["nav_sar"] is None
    assert "portfolio_decision_kpi_cash_vs_snapshot" not in findings
    assert "portfolio_decision_kpi_nav_vs_recomputed" not in findings
    assert "panel_cash_vs_snapshot" not in findings


@pytest.mark.parametrize("changes,reason", [
    ({"Date": dt.date(2026, 10, 9), "Time": None}, "future_snapshot"),
    ({"Date": dt.date(2026, 10, 9)}, "future_snapshot"),
    ({"Time": None}, "missing_or_invalid_timestamp"),
    ({"Time": "invalid"}, "missing_or_invalid_timestamp"),
    ({"Date": "notes: 2026-10-08"}, "missing_or_invalid_timestamp"),
    ({"Date": "2026-02-30"}, "missing_or_invalid_timestamp"),
    ({"Account": None}, "missing_account"),
    ({"Currency": None}, "missing_or_invalid_currency"),
    ({"Balance Type": None}, "balance_type_must_be_settled_cash"),
    ({"Balance Type": "available_funds", "Balance": 16_609.82},
     "balance_type_must_be_settled_cash"),
    ({"Balance": None}, "missing_or_invalid_native_balance"),
    ({"Balance": "NaN"}, "missing_or_invalid_native_balance"),
    ({"Balance": float("inf")}, "missing_or_invalid_native_balance"),
    ({"Balance": "1E9999999"}, "missing_or_invalid_native_balance"),
    ({"Balance": True}, "missing_or_invalid_native_balance"),
    ({"Balance": -1}, "missing_or_invalid_native_balance"),
    ({"FX to SAR": None}, "missing_or_invalid_fx_rate"),
    ({"FX to SAR": "Infinity"}, "missing_or_invalid_fx_rate"),
    ({"FX to SAR": 0}, "missing_or_invalid_fx_rate"),
    ({"FX As Of": None}, "missing_or_invalid_fx_timestamp"),
    ({"FX As Of": "2026-10-08"}, "missing_or_invalid_fx_timestamp"),
    ({"FX As Of": "2026-10-08T12:01:00Z"}, "future_fx_timestamp"),
    ({"FX As Of": "2026-10-08T11:31:00Z"}, "fx_timestamp_after_cash_snapshot"),
    ({"FX As Of": "2026-10-07T11:59:00Z"}, "stale_fx_timestamp"),
    ({"Balance SAR": 46_399.50}, "recorded_sar_disagrees_with_native_balance_and_fx"),
    ({"Balance SAR": "NaN"}, "invalid_recorded_sar_balance"),
    ({"Balance": "1E308", "FX to SAR": "1E308"}, "invalid_converted_sar_balance"),
])
def test_invalid_latest_cash_never_certifies_cash_or_nav(tmp_path, changes, reason):
    findings, metrics = _check(_write(tmp_path, _record(**changes)))
    _assert_uncertified(findings, metrics, reason)


def test_date_only_asof_does_not_get_a_midnight_timestamp(tmp_path):
    path = _write(tmp_path, _record(**{"As Of": "2026-10-08"}), extra_headers=("As Of",))
    findings, metrics = _check(path)
    _assert_uncertified(findings, metrics, "missing_or_invalid_timestamp")


@pytest.mark.parametrize("blank", [None, "", " \t "])
@pytest.mark.parametrize("source", ["timestamp", "date_time", "timestamp_only_date_time"])
def test_blank_timestamp_aliases_defer_to_the_next_available_source(tmp_path, blank, source):
    if source == "timestamp":
        changes = {"As Of": blank, "Timestamp": "2026-10-08T11:30:00Z"}
        extra_headers = ("As Of", "Timestamp")
    elif source == "date_time":
        changes = {"As Of": blank, "Timestamp": blank}
        extra_headers = ("As Of", "Timestamp")
    else:
        changes = {"Timestamp": blank}
        extra_headers = ("Timestamp",)
    findings, metrics = _check(_write(tmp_path, _record(**changes), extra_headers=extra_headers))
    assert findings["cash_snapshot_certification"]["status"] == "PASS"
    assert metrics["cash_snapshot"]["as_of_utc"] == "2026-10-08T11:30:00+00:00"
    assert metrics["cash_sar"] == pytest.approx(46_444.4642088, rel=0, abs=1e-9)


@pytest.mark.parametrize("header", ["As Of", "Timestamp"])
@pytest.mark.parametrize("invalid", ["invalid", "2026-10-08", 0, False, dt.date(2026, 10, 8)])
def test_populated_invalid_timestamp_alias_cannot_fall_through_to_valid_sources(tmp_path, header, invalid):
    changes = {"As Of": None, "Timestamp": "2026-10-08T11:30:00Z", header: invalid}
    path = _write(tmp_path, _record(**changes), extra_headers=("As Of", "Timestamp"))
    findings, metrics = _check(path)
    _assert_uncertified(findings, metrics, "missing_or_invalid_timestamp")


def test_blank_full_timestamp_alias_keeps_date_only_excel_time_fail_closed(tmp_path):
    path = _write(tmp_path, _record(**{"As Of": None, "Time": dt.date(2026, 10, 8)}),
                  extra_headers=("As Of",))
    findings, metrics = _check(path)
    _assert_uncertified(findings, metrics, "missing_or_invalid_timestamp")


@pytest.mark.parametrize("header,reason", [
    ("As Of", "missing_or_invalid_timestamp"),
    ("Timestamp", "missing_or_invalid_timestamp"),
    ("FX As Of", "missing_or_invalid_fx_timestamp"),
    ("Time", "missing_or_invalid_timestamp"),
])
def test_native_excel_date_only_cells_do_not_become_complete_timestamps(tmp_path, header, reason):
    extras = (header,) if header not in HEADERS else ()
    path = _write(tmp_path, _record(**{header: dt.date(2026, 10, 8)}), extra_headers=extras)
    wb = openpyxl.load_workbook(path)
    ws = wb["_Cash_Snapshot"]
    column = next(cell.column for cell in ws[1] if cell.value == header)
    ws.cell(2, column).number_format = "yyyy-mm-dd"
    wb.save(path)
    # openpyxl turns this native date into datetime(midnight); formatting is
    # the retained witness that no time was supplied.
    assert isinstance(audit.Book(path).rows("_Cash_Snapshot")[1][column - 1], dt.datetime)
    findings, metrics = _check(path)
    _assert_uncertified(findings, metrics, reason)


@pytest.mark.parametrize("header", ["As Of", "Timestamp", "FX As Of"])
def test_native_excel_complete_midnight_timestamp_remains_valid(tmp_path, header):
    extras = (header,) if header not in HEADERS else ()
    row = _record(**{header: dt.datetime(2026, 10, 8)})
    if header in ("As Of", "Timestamp"):
        row["FX As Of"] = "2026-10-07T20:59:00Z"
    path = _write(tmp_path, row, extra_headers=extras)
    wb = openpyxl.load_workbook(path)
    ws = wb["_Cash_Snapshot"]
    column = next(cell.column for cell in ws[1] if cell.value == header)
    ws.cell(2, column).number_format = "yyyy-mm-dd hh:mm:ss"
    wb.save(path)
    findings, metrics = _check(path)
    assert findings["cash_snapshot_certification"]["status"] == "PASS"
    assert metrics["cash_certified"] is True


def test_invalid_last_append_does_not_fall_back_to_valid_older_cash(tmp_path):
    earlier = _record(**{"Time": dt.time(14), "FX As Of": "2026-10-08T10:59:00Z"})
    latest = _record(**{"Date": dt.date(2026, 10, 9), "Time": None})
    findings, metrics = _check(_write(tmp_path, earlier, latest))
    _assert_uncertified(findings, metrics, "future_snapshot")
    assert metrics["cash_sar_raw"] == 46_444.46


def test_invalid_last_balance_does_not_disappear_from_the_record_selection(tmp_path):
    findings, metrics = _check(_write(tmp_path, _record(), _record(**{
        "Balance": "invalid", "Balance SAR": "invalid"})))
    _assert_uncertified(findings, metrics, "missing_or_invalid_native_balance")
    assert metrics["cash_sar_raw"] is None


def test_explicit_invalid_sar_native_balance_is_not_replaced_by_cached_sar(tmp_path):
    findings, metrics = _check(_write(tmp_path, _record(**{
        "Currency": "SAR", "Balance": "invalid"})))
    _assert_uncertified(findings, metrics, "missing_or_invalid_native_balance")


def test_legacy_sar_header_layout_requires_new_account_currency_and_balance_type(tmp_path):
    path = tmp_path / "legacy.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "_Cash_Snapshot"
    ws.append(["Date", "Time", "Balance SAR"])
    ws.append([dt.date(2026, 10, 8), dt.time(14, 30), 46_398.75])
    wb.save(path)
    findings, metrics = _check(path)
    _assert_uncertified(findings, metrics, "missing_account")
    assert "missing_or_invalid_currency" in metrics["cash_certification_errors"]
    assert "balance_type_must_be_settled_cash" in metrics["cash_certification_errors"]


def test_native_usd_cash_uses_recorded_fx_and_vintage(tmp_path):
    path = _write(tmp_path, _record())
    findings, metrics = _check(path)
    # Broker balance from the audit, with the workbook's recorded FX.
    assert metrics["cash_sar"] == pytest.approx(46_444.4642088, rel=0, abs=1e-9)
    assert metrics["cash_sar"] != 12_373.20 * audit.FX_TO_SAR["USD"]
    assert metrics["cash_sar_raw"] == 46_444.46
    assert metrics["cash_snapshot"]["fx_to_sar"] == "3.753634"
    assert metrics["cash_snapshot"]["fx_as_of_utc"] == "2026-10-08T11:25:00+00:00"
    assert metrics["cash_certified"] is True
    assert metrics["cash_certification_errors"] == []
    assert findings["cash_snapshot_certification"]["status"] == "PASS"
    assert findings["portfolio_decision_kpi_cash_vs_snapshot"]["status"] == "PASS"
    assert metrics["nav_sar"] == 46_444.46


@pytest.mark.parametrize("clock", [dt.time(0), 0, "00:00:00"])
def test_explicit_midnight_is_valid_and_sar_needs_no_market_fx_read(tmp_path, clock):
    row = _record(**{"Currency": "SAR", "Time": clock, "Balance": 46_444.46,
                     "FX to SAR": None, "FX As Of": None})
    findings, metrics = _check(_write(tmp_path, row))
    assert findings["cash_snapshot_certification"]["status"] == "PASS"
    assert metrics["cash_sar"] == 46_444.46
    assert metrics["cash_snapshot"]["as_of_utc"] == "2026-10-07T21:00:00+00:00"
    assert metrics["cash_snapshot"]["fx_to_sar"] == "1"


def test_stale_cash_and_explicit_maximum_age_policy(tmp_path):
    path = _write(tmp_path, _record(**{"Date": dt.date(2026, 10, 7), "Time": dt.time(14),
                                     "FX As Of": "2026-10-07T10:59:00Z"}))
    findings, metrics = _check(path)
    _assert_uncertified(findings, metrics, "stale_snapshot")
    findings, metrics = _check(path, cash_max_age_hours=48)
    assert findings["cash_snapshot_certification"]["status"] == "PASS"
    assert metrics["cash_certified"] is True


@pytest.mark.parametrize("maximum", [0, -1, float("inf"), float("nan")])
def test_invalid_age_policy_cannot_disable_certification_checks(tmp_path, maximum):
    findings, metrics = _check(_write(tmp_path, _record()), cash_max_age_hours=maximum)
    _assert_uncertified(findings, metrics, "invalid_cash_max_age_policy")


def test_snapshot_record_is_immutable_and_decimal_conversion_is_exact():
    row = _record()
    record = audit._cash_snapshot_record(dict(zip(HEADERS, range(len(HEADERS)))),
                                         [row.get(header) for header in HEADERS], 3)
    with pytest.raises(dataclasses.FrozenInstanceError):
        record.balance = Decimal("1")
    converted, errors = record.certify(NOW)
    assert converted == Decimal("46444.4642088")
    assert errors == ()


def test_real_cli_fails_future_cash_with_reviewable_reasons(tmp_path):
    path = tmp_path / "future.xlsx"
    expected = audit._build_synthetic(path, True, asof="2026-10-08")
    wb = openpyxl.load_workbook(path)
    ws = wb["_Cash_Snapshot"]
    ws.cell(3, 1, dt.datetime(2026, 10, 9))
    ws.cell(3, 2).value = None
    wb.save(path)
    result_path = tmp_path / "audit.json"
    cmd = [sys.executable, audit.__file__, str(path), "--expect",
           ",".join(f"{name}={count}" for name, count in expected.items()),
           "--now", NOW.isoformat(), "--asof", "2026-10-08", "--quiet",
           "--json", str(result_path), "--cash-max-age-hours", "24"]
    result = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
    assert result.returncode == 1, result.stderr
    payload = json.loads(result_path.read_text())
    assert payload["lanes"]["BOOK"]["rag"] == "RED"
    assert payload["metrics"]["book"]["cash_sar"] is None
    assert payload["metrics"]["book"]["nav_sar"] is None
    assert "future_snapshot" in payload["metrics"]["book"]["cash_certification_errors"]
    assert payload["params"]["cash_max_age_hours"] == 24
    assert any("cash_snapshot_certification" in fix for fix in payload["fix_list"])
    assert "BOOK=RED" in result.stdout
